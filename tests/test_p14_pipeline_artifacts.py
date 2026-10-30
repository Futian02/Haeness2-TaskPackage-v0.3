import csv
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
STAGE = ROOT / "work" / "data_staging" / "p14"


def read_jsonl(path):
    with path.open(encoding="utf-8-sig") as f:
        return [json.loads(line) for line in f if line.strip()]


class P14PipelineArtifactTests(unittest.TestCase):
    def test_page_coverage_matrix_is_full_census_and_unreviewed(self):
        path = ROOT / "docs" / "data_quality" / "P14_SOURCE_COVERAGE_MATRIX.csv"
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows), 1385)
        self.assertEqual(len({(x["source_pdf"], x["pdf_page_index"]) for x in rows}), 1385)
        self.assertTrue(all(x["human_visual_review"] == "False" for x in rows))

    def test_candidate_workbooks_preserve_teacher_four_sheet_templates(self):
        expected = [("条款摘录表_RBI_candidate.xlsx", 23, 472), ("条款摘录表_损伤机理_candidate.xlsx", 25, 82)]
        for name, columns, records in expected:
            wb = load_workbook(STAGE / name, read_only=True, data_only=False)
            self.assertEqual(len(wb.sheetnames), 4)
            self.assertEqual(wb.sheetnames, ["填写说明", "摘录表", "填写示例", "枚举值"])
            ws = wb["摘录表"]
            self.assertEqual(len(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))), columns)
            self.assertEqual(sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]), records)
            wb.close()

    def test_teacher_converter_candidate_is_field_identical_to_frozen_jsonl(self):
        candidate = read_jsonl(STAGE / "clauses_candidate.jsonl")
        official = read_jsonl(ROOT / "upstream" / "05_提交区" / "clauses.jsonl")
        self.assertEqual(len(candidate), 554)
        self.assertEqual(len({r["clause_id"] for r in candidate}), 554)
        self.assertEqual({r["clause_id"]: r for r in candidate}, {r["clause_id"]: r for r in official})

    def test_l2_registry_is_not_promoted_to_human_verified(self):
        rows = read_jsonl(STAGE / "candidates" / "p14_source_candidates.jsonl")
        self.assertEqual(len(rows), 568)
        self.assertEqual(len({r["candidate_id"] for r in rows}), 568)
        self.assertTrue(all(r.get("human_visual_review") is not True for r in rows))
        self.assertTrue(all(r.get("evidence_level") != "human_verified" for r in rows))

    def test_20_page_pilot_records_unscored_extraction_exactness(self):
        with (STAGE / "pilot" / "pilot_20_pages.csv").open(encoding="utf-8-sig", newline="") as f:
            pages = list(csv.DictReader(f))
        with (STAGE / "pilot" / "pilot_ocr_metrics.csv").open(encoding="utf-8-sig", newline="") as f:
            ocr = list(csv.DictReader(f))
        self.assertEqual(len(pages), 20)
        self.assertEqual(len(ocr), 20)
        self.assertTrue(all(int(p["candidate_count"]) == 0 for p in pages))
        self.assertTrue(all(p["ocr_exactness"] == "not_scored_against_human_transcription" for p in ocr))


if __name__ == "__main__":
    unittest.main()
