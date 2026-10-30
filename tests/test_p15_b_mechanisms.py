from __future__ import annotations
import csv, json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"work/data_staging/p15/B_DM"
DOCS=ROOT/"docs/data_quality"

def test_candidate_keeps_82_rows_25_columns_and_no_forged_reviewers():
    wb=load_workbook(OUT/"条款摘录表_损伤机理_P15_B_candidate.xlsx",data_only=True)
    assert wb.sheetnames==["填写说明","摘录表","填写示例","枚举值"]
    ws=wb["摘录表"]
    assert ws.max_column==25
    ids=[ws.cell(r,1).value for r in range(2,ws.max_row+1) if ws.cell(r,1).value]
    assert len(ids)==82 and len(set(ids))==82
    assert all(ws.cell(r,24).value in (None,"") for r in range(2,ws.max_row+1) if ws.cell(r,1).value)

def test_high_risk_condition_semantics_are_separated():
    ws=load_workbook(OUT/"条款摘录表_损伤机理_P15_B_candidate.xlsx",data_only=True)["摘录表"]
    byid={ws.cell(r,1).value.split("-")[-1]:r for r in range(2,ws.max_row+1) if ws.cell(r,1).value}
    r=byid["5.6"]
    assert ws.cell(r,17).value==177 and ws.cell(r,18).value==427
    assert "218℃～400℃" in ws.cell(r,22).value and "轻相介质" in ws.cell(r,22).value
    r=byid["5.23"]
    assert ws.cell(r,17).value is None and ws.cell(r,18).value is None
    assert "49℃～66℃" in ws.cell(r,22).value and "不能作为全机理" in ws.cell(r,22).value
    r=byid["6.7"]
    assert ws.cell(r,18).value is None and ws.cell(r,19).value is None and ws.cell(r,20).value is None
    assert "分压" in ws.cell(r,22).value and "0.0003 MPa" in ws.cell(r,22).value

def test_all_656_gn_slots_are_located_but_not_misrepresented_as_human_verified():
    with (OUT/"DM_FIELD_TRIAGE.csv").open(encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
    assert len(rows)==1312
    gn=[r for r in rows if r["field_letter"] in list("GHIJKLMN")]
    conditions=[r for r in rows if r["field_letter"] in list("OPQRSTUV")]
    assert len(gn)==656 and len(conditions)==656
    assert len({(r["clause_id"],r["field_letter"]) for r in rows})==1312
    assert all(r["reviewer_status"]=="NOT_HUMAN_VERIFIED" for r in rows)
    assert sum(r["ocr_heading_found"]=="True" for r in gn)==594
    assert all(r["source_bbox_or_table_cell"]=="NOT_CAPTURED" for r in rows)

def test_supplements_are_chapter_classified_and_not_mixed_into_dm():
    with (DOCS/"P15_B_SUPPLEMENT_CHAPTER_DISPOSITION.csv").open(encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
    assert len(rows)==84
    assert all(r["disposition"] and r["rationale"] for r in rows)
    assert all(r["direct_dm_candidate"]=="NO" for r in rows)
    index=[json.loads(x) for x in (OUT/"supplemental_index.jsonl").read_text(encoding="utf-8").splitlines()]
    assert index and all(r["raw_quote"]=="" and r["review_status"]=="REVIEW_NEEDED" for r in index)

def test_report_and_summary_are_blocked_until_human_review():
    s=json.loads((OUT/"run_summary.json").read_text(encoding="utf-8"))
    assert s["status"]=="REVIEW_BLOCKED" and s["human_review_completed"] is False
    assert s["reviewer_cells_blank"]==82
    assert s["direct_supplement_dm_candidates"]==0
    assert "REVIEW_BLOCKED" in (DOCS/"P15_B_MECHANISM_FIELD_REVIEW.md").read_text(encoding="utf-8")
