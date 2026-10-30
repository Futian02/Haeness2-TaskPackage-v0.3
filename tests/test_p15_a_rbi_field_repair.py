from __future__ import annotations
import csv, hashlib, json, unittest
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'work/data_staging/p15/A_RBI'
SOURCE=ROOT/'upstream/01_RBI标准/条款摘录表_RBI.xlsx'
CAND=OUT/'条款摘录表_RBI_P15_A_candidate.xlsx'
CHANGE=OUT/'P15_A_FIELD_CHANGELOG.csv'


def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def rows(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

class TestP15ARbiFieldRepair(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.src=load_workbook(SOURCE,read_only=False,data_only=True)
        cls.cand=load_workbook(CAND,read_only=False,data_only=True)
        cls.sws=cls.src['摘录表']; cls.cws=cls.cand['摘录表']
        cls.headers=[x.value for x in cls.sws[1]]
        cls.change=rows(CHANGE)

    def test_candidate_template_and_row_count(self):
        self.assertEqual(self.src.sheetnames,self.cand.sheetnames)
        self.assertEqual(23,self.cws.max_column)
        self.assertEqual(self.headers,[self.cws.cell(1,c).value for c in range(1,24)])
        nonempty=sum(bool(self.sws.cell(r,1).value) for r in range(2,self.sws.max_row+1))
        self.assertEqual(485,nonempty)
        self.assertEqual(nonempty,sum(bool(self.cws.cell(r,1).value) for r in range(2,self.cws.max_row+1)))

    def test_every_candidate_cell_diff_has_a_change_record_and_no_unlogged_diff(self):
        logged={(int(x['excel_row']),x['field']) for x in self.change}
        changed=set()
        for r in range(2,self.sws.max_row+1):
            for c,h in enumerate(self.headers,1):
                if self.sws.cell(r,c).value!=self.cws.cell(r,c).value:
                    changed.add((r,h))
        self.assertEqual(changed,logged)
        self.assertEqual(64,len(changed))

    def test_only_evidence_scoped_fields_changed(self):
        allowed=set()
        allowed.update((r,'其他适用条件') for r in range(2,27))
        allowed.update((r,'备注') for r in range(2,27))
        allowed.update((r,'原文摘录') for r in range(144,154))
        allowed.update((r,'其他适用条件') for r in (450,451,454,456))
        self.assertEqual(allowed,{(int(x['excel_row']),x['field']) for x in self.change})

    def test_no_extractor_or_reviewer_identity_touched_and_no_fake_signatures(self):
        for r in range(2,self.sws.max_row+1):
            self.assertEqual(self.sws.cell(r,21).value,self.cws.cell(r,21).value)
            self.assertEqual(self.sws.cell(r,22).value,self.cws.cell(r,22).value)
        self.assertTrue(all(x['review_required'] in ('TRUE_HUMAN_REQUIRED','TRUE_HUMAN_VISUAL_CONFIRMATION','TRUE') for x in self.change))
        self.assertEqual(0,sum(bool(self.cws.cell(r,22).value) for r in range(2,self.cws.max_row+1)))

    def test_table_review_queue_is_complete_and_unreviewed(self):
        x=rows(OUT/'P15_A_TABLE_REVIEW_COVERAGE.csv')
        self.assertEqual(221,len(x))
        self.assertEqual(221,len({q['excel_row'] for q in x}))
        self.assertTrue(all(q['reviewer_signature']=='EMPTY' for q in x))
        self.assertTrue(all(q['review_status'] for q in x))

    def test_field_triage_is_complete(self):
        x=rows(OUT/'RBI_FIELD_TRIAGE.csv')
        self.assertEqual(485*23,len(x))
        self.assertEqual(485*23,len({(q['excel_row'],q['field']) for q in x}))
        classes={'REQUIRED','CONDITIONAL_SUPPORTED','NOT_APPLICABLE','SOURCE_UNSPECIFIED','REVIEW_NEEDED','PROPOSED_CORRECTION'}
        self.assertTrue(set(q['field_class'] for q in x)<=classes)

    def test_source_sha_unchanged_from_frozen_p14_snapshot(self):
        self.assertEqual('e43f7b0d3381a38fa648c2505e27d1c87893af5aea89aa119097dbfc49c6b7bf',sha(SOURCE))
        self.assertNotEqual(sha(SOURCE),sha(CAND))

    def test_candidate_f_quotes_respect_300_char_contract(self):
        for r in range(2,self.cws.max_row+1):
            if self.cws.cell(r,1).value:
                self.assertLessEqual(len(str(self.cws.cell(r,6).value or '')),300,r)

    def test_visual_table_excerpt_values_units_and_bounds(self):
        expected={
            144:'表2 面积后果等级划分\n单位为平方米\nA\nCA≤9.29',
            145:'表2 面积后果等级划分\n单位为平方米\nB\n9.29<CA≤92.9',
            146:'表2 面积后果等级划分\n单位为平方米\nC\n92.9<CA≤279',
            147:'表2 面积后果等级划分\n单位为平方米\nD\n279<CA≤929',
            148:'表2 面积后果等级划分\n单位为平方米\nE\nCA>929',
            149:'表3 热交换器管束的经济后果等级划分\n单位为元\nA\nCOF≤50 000',
            150:'表3 热交换器管束的经济后果等级划分\n单位为元\nB\n50 000<COF≤250 000',
            151:'表3 热交换器管束的经济后果等级划分\n单位为元\nC\n250 000<COF≤750 000',
            152:'表3 热交换器管束的经济后果等级划分\n单位为元\nD\n750 000<COF≤5 000 000',
            153:'表3 热交换器管束的经济后果等级划分\n单位为元\nE\nCOF>5 000 000',
        }
        for r,v in expected.items(): self.assertEqual(v,self.cws.cell(r,6).value)

    def test_matrix_editorial_t_is_blank_and_trace_is_preserved_in_w(self):
        for r in range(2,27):
            self.assertIsNone(self.cws.cell(r,20).value)
            self.assertIn('图1风险矩阵共25格',self.cws.cell(r,23).value)
            self.assertIn('图1在正文第14页（PDF index 13）',self.cws.cell(r,23).value)
            self.assertNotIn('图1在正文第12页',self.cws.cell(r,23).value)

    def test_thirteen_unsigned_drafts_remain_unsigned(self):
        n=0
        for r in range(2,self.cws.max_row+1):
            if self.cws.cell(r,2).value=='GB/T 26610.1-2022':
                n+=1
                self.assertFalse(self.cws.cell(r,21).value)
                self.assertFalse(self.cws.cell(r,22).value)
        self.assertEqual(13,n)

if __name__=='__main__': unittest.main()

