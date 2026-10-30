# -*- coding: utf-8 -*-
"""
把 26610.3 的 11 条「一条含多个 a)~e) 赋值选项」的因子条款，按选项拆成多行。

原则：
1. 引导句（第一个选项之前的那段）保留在每个选项行前，与 EF 表(5.2.1-a/b/c)、
   DF 表(5.2.2-a~j) 的既有拆行体例一致。
2. 每段的文字**全部来自原文字符串本身**，不做任何改写、增删、标点替换。
3. 原文已逐页读图核对（印刷页 4~6 / 物理页 6~8），与库中现有 content_raw 逐字一致。
4. 拆行前逐条校验：切出的每段必须能在原文中找到（is 关系，非 in 关系），
   引导句必须等于原文的前缀。
"""
import openpyxl, re, sys

XLSX = '01_RBI标准/条款摘录表_RBI.xlsx'
TARGET_STD = 'GB/T 26610.3-2014'
# 目标 clause_no（11 条，各含多个 a)~e) 选项）
TARGET_NOS = ['5.2.3.1', '5.2.3.2', '5.2.3.3',
              '5.2.4.1', '5.2.4.2', '5.2.4.3',
              '5.2.5.1', '5.2.5.2', '5.2.5.3',
              '5.2.6.1', '5.2.6.2']

wb = openpyxl.load_workbook(XLSX)
ws = wb['摘录表']
nh = [c.value for c in ws[1]]
idx = {h: i + 1 for i, h in enumerate(nh)}

rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]

by_pos = {}
for r in rows:
    std = ws.cell(r, idx['standard_code']).value
    no = str(ws.cell(r, idx['clause_no']).value)
    if std != TARGET_STD or no not in TARGET_NOS:
        continue
    d = {h: ws.cell(r, idx[h]).value for h in nh}
    orig = d['原文摘录']
    # 切分：引导句 + 各选项
    parts = re.split(r'(?=[a-e]\)\s)', orig)
    lead = parts[0]
    opts = [p for p in parts[1:] if p.strip()]
    if len(opts) < 2:
        print('✗ %s 未切出多个选项（只有 %d 个）' % (no, len(opts)))
        sys.exit(1)
    # 校验：每段必须在原文中，且拼接后与原文完全一致（除分隔符外不改一字）
    rebuilt = lead + ''.join(opts)
    if rebuilt != orig:
        print('✗ %s 重组后与原文不一致' % no)
        sys.exit(1)
    lst = []
    for p in opts:
        letter = p[0]
        nd = dict(d)
        nd['clause_id'] = '%s-%s' % (d['clause_id'], letter)
        nd['原文摘录'] = lead + p
        m = re.search(r'取\s*([A-Z]+\d+)\s*=\s*(-?\d+(?:\.\d+)?)', p)
        code, val = (m.group(1), m.group(2)) if m else ('', '')
        nd['其他适用条件'] = '%s 第%s)项：%s=%s' % (no, letter, code, val)
        lst.append(nd)
    by_pos[r] = lst
    print('  %-9s 1 条 -> %d 条  选项: %s' % (
        no, len(opts), [p[0] for p in opts]))

if len(by_pos) != len(TARGET_NOS):
    print('✗ 命中 %d 条，期望 %d 条' % (len(by_pos), len(TARGET_NOS)))
    sys.exit(1)

ordered = []
for r in rows:
    if r in by_pos:
        ordered.extend(by_pos[r])
    else:
        ordered.append({h: ws.cell(r, idx[h]).value for h in nh})

for r in range(2, ws.max_row + 1):
    for c in range(1, len(nh) + 1):
        ws.cell(r, c).value = None

for i, d in enumerate(ordered):
    r = 2 + i
    for c in range(1, len(nh) + 1):
        ws.cell(r, c).value = d.get(nh[c - 1])

wb.save(XLSX)
print('写回完成：总行数 %d（原 %d，新增 %d）' % (len(ordered), len(rows), len(ordered) - len(rows)))
