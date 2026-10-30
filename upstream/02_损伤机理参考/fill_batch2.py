# -*- coding: utf-8 -*-
"""合并批次2 的 4 份子代理 JSON，按条款号排序后写入损伤机理表（从第33行起）。
图片由子代理在隔离上下文转录，本脚本只负责合并落盘，不读图。
"""
import json
import copy
import os
import glob

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "条款摘录表_损伤机理.xlsx")
SHEET = "摘录表"

COLS = ["clause_id", "standard_code", "clause_no", "page", "title", "原文摘录",
        "损伤描述及机理", "损伤形态", "受影响的材料", "主要影响因素", "易发生部位",
        "主要预防措施", "检测监测方法", "相关或伴随损伤", "适用材质", "适用介质",
        "温度下限", "温度上限", "压力下限", "压力上限", "适用设备类型",
        "其他适用条件", "抽取人", "复核人", "备注"]

JSON_FILES = ["batch2_partA.json", "batch2_partB.json",
              "batch2_partC.json", "batch2_partD.json"]


def clause_key(no):
    parts = []
    for p in str(no).split("."):
        try:
            parts.append(int(p))
        except ValueError:
            parts.append(0)
    return tuple(parts)


def main():
    rows = []
    uncertain_all = []
    for jf in JSON_FILES:
        path = os.path.join(HERE, jf)
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for d in data:
            rows.append(d)
            unc = d.get("_uncertain") or []
            for u in unc:
                uncertain_all.append(f"[{d.get('clause_no','?')}] {u}")

    rows.sort(key=lambda d: clause_key(d.get("clause_no", "")))

    wb = load_workbook(XLSX)
    ws = wb[SHEET]

    # 找到最后有 clause_id 的行
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value not in (None, ""):
            last = r
    start = last + 1

    n = len(rows)
    for i, d in enumerate(rows):
        r = start + i
        src = ws.cell(last, 1).row  # 样式源行 = last
        for c, key in enumerate(COLS, start=1):
            s_cell = ws.cell(last, c)
            d_cell = ws.cell(r, c)
            # 复制样式
            d_cell.font = copy.copy(s_cell.font)
            d_cell.border = copy.copy(s_cell.border)
            d_cell.fill = copy.copy(s_cell.fill)
            d_cell.alignment = copy.copy(s_cell.alignment)
            val = d.get(key)
            # None / 空串 → 留空
            d_cell.value = val if val not in (None, "") else None
            if key == "clause_no":
                d_cell.number_format = "@"
        # 强制规范抽取人/复核人/备注
        ws.cell(r, COLS.index("抽取人") + 1).value = "AI试抽(WorkBuddy)"
        ws.cell(r, COLS.index("复核人") + 1).value = ""
        ws.cell(r, COLS.index("备注") + 1).value = "REMARK"

    wb.save(XLSX)

    print(f"写入 {n} 条，行 {start}~{start + n - 1}")
    print(f"条款号顺序: {', '.join(d.get('clause_no','') for d in rows)}")
    print(f"子代理标注的不确定项共 {len(uncertain_all)} 条:")
    for u in uncertain_all:
        print("  -", u)


if __name__ == "__main__":
    main()
