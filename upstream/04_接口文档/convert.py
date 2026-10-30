#!/usr/bin/env python3
"""把两份条款摘录表转换成 clauses.jsonl（v0.3）。

依赖：python -m pip install openpyxl

在任务包根目录运行：
  python 04_接口文档/convert.py
  python 04_接口文档/convert.py --inputs 01_RBI标准/条款摘录表_RBI.xlsx --output 05_提交区/clauses.jsonl

默认读取：
  01_RBI标准/条款摘录表_RBI.xlsx
  02_损伤机理参考/条款摘录表_损伤机理.xlsx
只读每个文件的「摘录表」页，「填写示例」页不参与转换。
默认输出 05_提交区/clauses.jsonl（UTF-8）。转换完成后请接着运行 validate.py。

本脚本做的检查：
  表头是否被改动；必填列；clause_id 是否由 standard_code 与 clause_no 拼成；
  主键重复；原文摘录 ≤300 字；页码与温压是否为数字；
  枚举列只允许 enums.json 的规范值（写成别名时自动改为规范值并提示）；
  GB/T 26610.3-2014 的风险等级按 risk_level_mapping.xlsx 核对，统一等级留空时自动补上并提示；
  GB/T 30579-2022 的损伤机理条款，title 必须与附录A 索引中该条的名称一致；
  损伤模式常见的错误叫法。
有任何错误时不写输出文件，返回退出码 1。报错请回到 xlsx 修改，不要手改 jsonl。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少依赖，请运行：python -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(2)

BASE_DIR = Path(__file__).resolve().parent
ROOT = BASE_DIR.parent
DEFAULT_INPUTS = [ROOT / "01_RBI标准" / "条款摘录表_RBI.xlsx", ROOT / "02_损伤机理参考" / "条款摘录表_损伤机理.xlsx"]
DEFAULT_OUTPUT = ROOT / "05_提交区" / "clauses.jsonl"
SHEET = "摘录表"

STANDARD_NAMES = {
    "GB/T 30579-2022": "承压设备损伤模式识别",
    "GB/T 26610.1-2022": "承压设备系统基于风险的检验实施导则 第1部分：基本要求和实施程序",
    "GB/T 26610.2-2022": "承压设备系统基于风险的检验实施导则 第2部分：基于风险的检验策略",
    "GB/T 26610.3-2014": "承压设备系统基于风险的检验实施导则 第3部分：风险的定性分析方法",
    "GB/T 26610.4-2022": "承压设备系统基于风险的检验实施导则 第4部分：失效可能性定量分析方法",
    "GB/T 26610.5-2022": "承压设备系统基于风险的检验实施导则 第5部分：失效后果定量分析方法",
    "TSG 21-2016": "固定式压力容器安全技术监察规程",
    "GB/T 21447-2018": "钢质管道外腐蚀控制规范",
    "GB/T 150.1-2024": "压力容器 第1部分：通用要求",
    "GB/T 150.3-2024": "压力容器 第3部分：设计",
    "GB/T 150.4-2024": "压力容器 第4部分：制造、检验和验收",
    "GB/T 151-2026": "热交换器",
}

KEY_COLS = ["clause_id", "standard_code", "clause_no", "page", "title", "原文摘录"]
MECH_COLS = ["损伤描述及机理", "损伤形态", "受影响的材料", "主要影响因素", "易发生部位", "主要预防措施", "检测监测方法", "相关或伴随损伤"]
RBI_COLS = ["失效可能性等级", "失效后果等级", "风险等级(原始)", "风险等级(统一)", "检验周期(年)", "推荐检验方法"]
COND_COLS = ["适用材质", "适用介质", "温度下限", "温度上限", "压力下限", "压力上限", "适用设备类型", "其他适用条件"]
META_COLS = ["抽取人", "复核人", "备注"]
HEADERS = {
    "damage_mechanism": KEY_COLS + MECH_COLS + COND_COLS + META_COLS,
    "rbi": KEY_COLS + RBI_COLS + COND_COLS + META_COLS,
}
MECH_FIELDS = {
    "损伤形态": "morphology", "受影响的材料": "affected_materials", "主要影响因素": "influencing_factors",
    "易发生部位": "susceptible_locations", "主要预防措施": "prevention", "检测监测方法": "detection_methods",
    "相关或伴随损伤": "related_damages",
}
WRONG_NAMES = {"湿硫化氢损伤": "湿硫化氢破坏", "氯化铵垢下腐蚀": "氯化铵腐蚀", "高温硫腐蚀": "高温硫化物腐蚀",
               "硫氢化铵腐蚀": "酸性水腐蚀(碱式酸性水)", "铵盐腐蚀": "氯化铵腐蚀 或 酸性水腐蚀(碱式酸性水)",
               "保温层下腐蚀": "大气腐蚀(有隔热层)"}

# GB/T 30579-2022 附录A 表A.1~表A.5 损伤模式索引:章条号 -> 标准名称(逐条核对原文)
ANNEX_A = {
    "5.1": "盐酸腐蚀", "5.2": "硫酸腐蚀", "5.3": "氢氟酸腐蚀", "5.4": "磷酸腐蚀", "5.5": "二氧化碳腐蚀",
    "5.6": "环烷酸腐蚀", "5.7": "苯酚腐蚀", "5.8": "低分子有机酸腐蚀", "5.9": "高温氧化",
    "5.10": "大气腐蚀(无隔热层)", "5.11": "大气腐蚀(有隔热层)", "5.12": "冷却水腐蚀", "5.13": "土壤腐蚀",
    "5.14": "微生物腐蚀", "5.15": "锅炉冷凝水腐蚀", "5.16": "碱腐蚀", "5.17": "燃灰腐蚀", "5.18": "烟气露点腐蚀",
    "5.19": "氯化铵腐蚀", "5.20": "胺腐蚀", "5.21": "高温硫化物腐蚀(无氢气环境)", "5.22": "高温硫化物腐蚀(氢气环境)",
    "5.23": "酸性水腐蚀(碱式酸性水)", "5.24": "酸性水腐蚀(酸式酸性水)", "5.25": "甲铵腐蚀", "5.26": "电偶腐蚀",
    "5.27": "含盐水腐蚀", "5.28": "含氧工艺水腐蚀", "5.29": "浓差电池腐蚀",
    "6.1": "氯化物应力腐蚀开裂", "6.2": "碳酸盐应力腐蚀开裂", "6.3": "硝酸盐应力腐蚀开裂", "6.4": "碱应力腐蚀开裂",
    "6.5": "氨应力腐蚀开裂", "6.6": "胺应力腐蚀开裂", "6.7": "湿硫化氢破坏", "6.8": "氢氟酸致氢应力开裂",
    "6.9": "氢氰酸致氢应力开裂", "6.10": "氢脆", "6.11": "高温水应力腐蚀开裂", "6.12": "连多硫酸应力腐蚀开裂",
    "6.13": "液体金属脆断", "6.14": "乙醇应力腐蚀开裂", "6.15": "硫酸盐应力腐蚀开裂", "6.16": "氢氟酸应力腐蚀开裂",
    "7.1": "晶粒长大", "7.2": "渗氮", "7.3": "球化", "7.4": "石墨化", "7.5": "渗碳", "7.6": "脱碳", "7.7": "金属粉化",
    "7.8": "σ相脆化", "7.9": "475℃脆化", "7.10": "回火脆化", "7.11": "辐照脆化", "7.12": "钛氢化", "7.13": "再热裂纹",
    "7.14": "脱金属腐蚀", "7.15": "敏化-晶间腐蚀", "7.16": "金属热老化",
    "8.1": "机械疲劳", "8.2": "热疲劳", "8.3": "振动疲劳", "8.4": "接触疲劳", "8.5": "机械磨损", "8.6": "冲刷",
    "8.7": "汽蚀", "8.8": "过载", "8.9": "热冲击", "8.10": "蠕变", "8.11": "应变时效",
    "9.1": "高温氢腐蚀", "9.2": "腐蚀疲劳", "9.3": "冲蚀", "9.4": "蒸汽阻滞", "9.5": "低温脆断", "9.6": "过热",
    "9.7": "耐火材料退化", "9.8": "铸铁石墨化腐蚀", "9.9": "微动腐蚀", "9.10": "高含氧气体促发的燃烧、爆炸",
}
TERM_UNIFIED = {"低风险": "L1", "中风险": "L2", "中高风险": "L3", "高风险": "L4"}


def norm_name(v: str) -> str:
    """比较损伤模式名称时统一全角/半角括号与空格。"""
    return re.sub(r"\s", "", v).replace("（", "(").replace("）", ")")
SPLIT = re.compile(r"[;；\n]+")


class Report:
    def __init__(self):
        self.errors = 0
        self.warnings = 0

    def error(self, where: str, msg: str) -> None:
        self.errors += 1
        print(f"[错误] {where}：{msg}", file=sys.stderr)

    def warn(self, where: str, msg: str) -> None:
        self.warnings += 1
        print(f"[提示] {where}：{msg}")


def load_enum_lookup() -> dict[str, dict[str, str]]:
    data = json.loads((BASE_DIR / "enums.json").read_text(encoding="utf-8-sig"))
    lookup = {}
    for table in ("material", "medium", "equipment_type"):
        m = {}
        for canonical, entry in data[table].items():
            m[canonical] = canonical
            for alias in entry.get("aliases", []):
                m[alias] = canonical
        lookup[table] = m
    return lookup


def load_risk_mapping() -> dict[tuple[str, str], tuple[str, str]]:
    path = BASE_DIR / "risk_level_mapping.xlsx"
    if not path.exists():
        return {}
    ws = load_workbook(path, read_only=True, data_only=True)["GBT26610.3-2014"]
    mapping = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            mapping[(str(row[0]).strip(), str(row[1]).strip().upper())] = (str(row[2]).strip(), str(row[3]).strip())
    return mapping


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def split_list(v) -> list[str]:
    return [x.strip() for x in SPLIT.split(text(v)) if x.strip()]


def number(v, where: str, rep: Report, *, integer: bool = False):
    if v is None or text(v) == "":
        return None
    try:
        n = float(text(v))
    except ValueError:
        rep.error(where, f"{text(v)!r} 不是数字")
        return None
    if integer:
        if not n.is_integer():
            rep.error(where, f"{text(v)!r} 必须是整数")
            return None
        return int(n)
    return int(n) if n.is_integer() else n


def enum_list(v, table: str, lookup, where: str, rep: Report) -> list[str]:
    out = []
    for item in split_list(v):
        canonical = lookup[table].get(item)
        if canonical is None:
            rep.error(where, f"{item!r} 不在 enums.json 的 {table} 表中；原文材料或介质超出枚举时写进「其他适用条件」")
            continue
        if canonical != item:
            rep.warn(where, f"别名 {item!r} 已改为规范值 {canonical!r}，建议在 xlsx 里也改掉")
        if canonical not in out:
            out.append(canonical)
    return out


def interval(lo, hi, where: str, rep: Report) -> list:
    a, b = number(lo, where + "/下限", rep), number(hi, where + "/上限", rep)
    if a is not None and b is not None and a > b:
        rep.error(where, f"下限 {a} 大于上限 {b}")
    return [a, b]


def check_names(values: list[str], where: str, rep: Report) -> None:
    for v in values:
        for wrong, right in WRONG_NAMES.items():
            if wrong in v:
                rep.error(where, f"{wrong!r} 不是 GB/T 30579-2022 的叫法，应为 {right!r}")


def convert_row(kind: str, row: dict, where: str, lookup, risk_map, rep: Report) -> dict | None:
    for col in ("clause_id", "standard_code", "clause_no", "page", "title", "原文摘录", "抽取人"):
        if text(row.get(col)) == "":
            rep.error(where, f"必填列「{col}」为空")
    if isinstance(row.get("clause_no"), float):
        rep.error(where, "clause_no 被 Excel 存成了数字（5.10 会变成 5.1），请把该列设为文本格式后重新输入")
    code, clause_no, clause_id = text(row.get("standard_code")), text(row.get("clause_no")), text(row.get("clause_id"))
    if code and code not in STANDARD_NAMES:
        rep.error(where, f"standard_code {code!r} 不在本任务包的标准清单里")
    if code and clause_no and clause_id:
        prefix = re.sub(r"[/\s]", "", code) + "-" + clause_no
        if not re.fullmatch(re.escape(prefix) + r"(-[a-z]{1,2})?", clause_id):
            rep.error(where, f"clause_id {clause_id!r} 应为 {prefix!r}，一个条款拆多行时加 -a/-b/-c")
    content = text(row.get("原文摘录"))
    if len(content) > 300:
        rep.error(where, f"原文摘录 {len(content)} 字，超过 300 字，请截取与判据相关的部分")
    title = text(row.get("title"))
    check_names([title], where + "/title", rep)
    if kind == "damage_mechanism" and code == "GB/T 30579-2022" and clause_no:
        top = ".".join(clause_no.split(".")[:2])
        expected = ANNEX_A.get(top)
        if expected is None:
            rep.error(where, f"GB/T 30579-2022 附录A 里没有第 {top} 条损伤模式，检查条款号")
        elif norm_name(title) != norm_name(expected):
            rep.error(where, f"第 {top} 条在 GB/T 30579-2022 里叫 {expected!r}，title 写的是 {title!r}")

    record = {
        "clause_id": clause_id,
        "standard_code": code,
        "standard_name": STANDARD_NAMES.get(code, ""),
        "clause_no": clause_no,
        "page": number(row.get("page"), where + "/page", rep, integer=True),
        "knowledge_type": kind,
        "title": title,
        "content_raw": content,
        "applicable_conditions": {
            "equipment_type": enum_list(row.get("适用设备类型"), "equipment_type", lookup, where + "/适用设备类型", rep),
            "material": enum_list(row.get("适用材质"), "material", lookup, where + "/适用材质", rep),
            "medium": enum_list(row.get("适用介质"), "medium", lookup, where + "/适用介质", rep),
            "temperature_c": interval(row.get("温度下限"), row.get("温度上限"), where + "/温度", rep),
            "pressure_mpa": interval(row.get("压力下限"), row.get("压力上限"), where + "/压力", rep),
            "other": text(row.get("其他适用条件")),
        },
        "mechanism_detail": {"description": "", "morphology": [], "affected_materials": [], "influencing_factors": [],
                             "susceptible_locations": [], "prevention": [], "detection_methods": [], "related_damages": []},
        "rbi_detail": {"likelihood": None, "consequence": None, "risk_level": None, "risk_level_unified": None,
                       "inspection_interval_year": None, "inspection_methods": []},
        "extractor": text(row.get("抽取人")),
        "verified_by": text(row.get("复核人")),
        "version": "v1",
    }

    if kind == "damage_mechanism":
        detail = record["mechanism_detail"]
        detail["description"] = text(row.get("损伤描述及机理"))
        if not detail["description"]:
            rep.error(where, "损伤机理条款的「损伤描述及机理」不能为空")
        for col, key in MECH_FIELDS.items():
            detail[key] = split_list(row.get(col))
        check_names(detail["related_damages"], where + "/相关或伴随损伤", rep)
    else:
        rbi = record["rbi_detail"]
        rbi["likelihood"] = text(row.get("失效可能性等级")) or None
        rbi["consequence"] = (text(row.get("失效后果等级")) or "").upper() or None
        rbi["risk_level"] = text(row.get("风险等级(原始)")) or None
        rbi["risk_level_unified"] = text(row.get("风险等级(统一)")).upper() or None
        rbi["inspection_interval_year"] = number(row.get("检验周期(年)"), where + "/检验周期(年)", rep)
        rbi["inspection_methods"] = split_list(row.get("推荐检验方法"))
        if rbi["risk_level_unified"] and rbi["risk_level_unified"] not in {"L1", "L2", "L3", "L4", "L5"}:
            rep.error(where, f"风险等级(统一) {rbi['risk_level_unified']!r} 只能是 L1～L5")
        if rbi["risk_level_unified"] == "L5":
            rep.error(where, "本任务包涉及的标准不会映射到 L5，请按 risk_level_mapping.xlsx 重查")
        matrix_row = code == "GB/T 26610.3-2014" and rbi["likelihood"] and rbi["consequence"]
        if not matrix_row and code.startswith("GB/T 26610") and rbi["risk_level"]:
            uni = TERM_UNIFIED.get(rbi["risk_level"])
            if uni is None:
                if rbi["risk_level_unified"]:
                    rep.error(where, f"原文用语 {rbi['risk_level']!r} 不在映射表「按用语对应」页，统一等级应留空，在备注写明")
            elif rbi["risk_level_unified"] and rbi["risk_level_unified"] != uni:
                rep.error(where, f"原文 {rbi['risk_level']!r} 按映射表应为 {uni}，表里写的是 {rbi['risk_level_unified']}")
            elif not rbi["risk_level_unified"]:
                rbi["risk_level_unified"] = uni
                rep.warn(where, f"风险等级(统一) 为空，已按用语补为 {uni}")
        if code == "GB/T 26610.3-2014" and rbi["likelihood"] and rbi["consequence"]:
            expected = risk_map.get((rbi["likelihood"], rbi["consequence"]))
            if expected is None:
                rep.error(where, f"映射表里没有 可能性{rbi['likelihood']}/后果{rbi['consequence']} 这一格")
            else:
                raw, uni = expected
                if rbi["risk_level"] and rbi["risk_level"] != raw:
                    rep.error(where, f"按标准图1，{rbi['likelihood']}{rbi['consequence']} 应为 {raw!r}，表里写的是 {rbi['risk_level']!r}")
                if rbi["risk_level_unified"] and rbi["risk_level_unified"] != uni:
                    rep.error(where, f"按映射表应为 {uni}，表里写的是 {rbi['risk_level_unified']}")
                if not rbi["risk_level_unified"]:
                    rbi["risk_level_unified"] = uni
                    rep.warn(where, f"风险等级(统一) 为空，已按映射表补为 {uni}")
    return record


def read_table(path: Path, lookup, risk_map, rep: Report) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        rep.error(str(path), f"找不到「{SHEET}」页")
        return []
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [text(h) for h in next(rows, [])]
    while header and header[-1] == "":
        header.pop()
    kind = next((k for k, cols in HEADERS.items() if header == cols), None)
    if kind is None:
        rep.error(str(path), "表头与 v0.3 模板不一致（列名、顺序都不能改）。请用任务包里的模板重新粘贴数据")
        return []
    records = []
    for i, values in enumerate(rows, start=2):
        row = dict(zip(header, values))
        if all(text(v) == "" for v in row.values()):
            continue
        rec = convert_row(kind, row, f"{path.name} 第{i}行", lookup, risk_map, rep)
        if rec:
            records.append((f"{path.name} 第{i}行", rec))
    return records


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="把条款摘录表（xlsx）转换成 clauses.jsonl")
    parser.add_argument("--inputs", type=Path, nargs="+", default=DEFAULT_INPUTS, help="摘录表 xlsx，可给多个")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="输出 jsonl 路径")
    args = parser.parse_args(argv)

    rep = Report()
    lookup, risk_map = load_enum_lookup(), load_risk_mapping()
    if not risk_map:
        rep.warn("risk_level_mapping.xlsx", "未找到映射表，跳过风险等级核对")
    all_records = []
    for path in args.inputs:
        if not path.exists():
            rep.error(str(path), "文件不存在")
            continue
        all_records += read_table(path, lookup, risk_map, rep)

    seen = {}
    for where, rec in all_records:
        cid = rec["clause_id"]
        if cid in seen:
            rep.error(where, f"clause_id {cid!r} 与 {seen[cid]} 重复")
        else:
            seen[cid] = where

    counts = {"damage_mechanism": 0, "rbi": 0}
    for _, rec in all_records:
        counts[rec["knowledge_type"]] += 1
    print(f"读取条款 {len(all_records)} 条：损伤机理 {counts['damage_mechanism']} 条，RBI {counts['rbi']} 条；"
          f"错误 {rep.errors} 处，提示 {rep.warnings} 处")
    if rep.errors:
        print("有错误，未写出 clauses.jsonl。请回到 xlsx 修改后重跑。", file=sys.stderr)
        return 1
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as f:
        for _, rec in all_records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(f"已写出 {args.output}，下一步：python 04_接口文档/validate.py")
    return 0


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    raise SystemExit(main())
