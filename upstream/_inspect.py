#!/usr/bin/env python3
# 临时检查脚本：核实现状
from openpyxl import load_workbook

def inspect(path, label):
    print(f"\n===== {label} : {path} =====")
    wb = load_workbook(path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    ws = wb["摘录表"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    print("ncols:", len(header))
    data = [r for r in rows[1:] if any(c is not None and str(c).strip()!="" for c in r)]
    print("data rows:", len(data))
    # 关键列索引
    idx = {h:i for i,h in enumerate(header)}
    print("header:", header)
    for r in data:
        cid = r[idx.get("clause_id")] if "clause_id" in idx else ""
        code = r[idx.get("standard_code")] if "standard_code" in idx else ""
        cno = r[idx.get("clause_no")] if "clause_no" in idx else ""
        title = r[idx.get("title")] if "title" in idx else ""
        print(f"  {cid} | {code} | {cno} | {title}")

inspect(r"D:\工业互联网\Haeness2\任务包v0.3\01_RBI标准\条款摘录表_RBI.xlsx", "RBI表")
inspect(r"D:\工业互联网\Haeness2\任务包v0.3\02_损伤机理参考\条款摘录表_损伤机理.xlsx", "损伤机理表")

# 检查损伤机理表是否含14条
targets = ["5.1","5.19","5.23","5.24","6.7","5.21","5.22","5.6","9.3","6.1","5.11","5.5","6.4","6.12"]
wb = load_workbook(r"D:\工业互联网\Haeness2\任务包v0.3\02_损伤机理参考\条款摘录表_损伤机理.xlsx", read_only=True, data_only=True)
ws = wb["摘录表"]
rows = list(ws.iter_rows(values_only=True))
header = [str(h).strip() if h is not None else "" for h in rows[0]]
idx = {h:i for i,h in enumerate(header)}
cno_i = idx["clause_no"]
found = set()
for r in rows[1:]:
    if r[cno_i] is not None:
        v = str(r[cno_i]).strip()
        top = ".".join(v.split(".")[:2])
        if top in targets:
            found.add(top)
print("\n14条常减压损伤模式在损伤机理表中命中:", sorted(found, key=lambda x:float(x)))
print("未命中:", [t for t in targets if t not in found])
