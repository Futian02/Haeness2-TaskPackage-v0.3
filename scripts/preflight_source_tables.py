"""Independent, READ-ONLY teacher-XLSX preflight for known convert.py blind spots.

Run: python -B -m scripts.preflight_source_tables
Reports cell coordinates and issue codes, never writes/mutates teacher assets.
Nonfatal review warnings are not equivalent to verified misconduct.
"""
from __future__ import annotations

import importlib.util
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

UPSTREAM = Path(__file__).resolve().parents[1] / "upstream"
CONVERT = UPSTREAM / "04_接口文档" / "convert.py"
TABLES = [
    ("rbi", UPSTREAM / "01_RBI标准" / "条款摘录表_RBI.xlsx"),
    ("damage_mechanism", UPSTREAM / "02_损伤机理参考" / "条款摘录表_损伤机理.xlsx"),
]
ENUM_COLS = ("适用材质", "适用介质", "适用设备类型")
LIST_COLS = ("损伤形态", "受影响的材料", "主要影响因素",
             "易发生部位", "主要预防措施", "检测监测方法", "相关或伴随损伤",
             "推荐检验方法")
SUSPICIOUS_SEPARATORS = ("、", ",")


def check() -> dict:
    spec = importlib.util.spec_from_file_location("official_convert_preflight", CONVERT)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot import official convert.py")
    conv = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(conv)
    problems = []
    inspected = Counter()
    summaries = []
    for kind, path in TABLES:
        if not path.is_file():
            problems.append(("fatal", str(path), "-", "MISSING_SOURCE"))
            continue
        workbook = load_workbook(path, read_only=True, data_only=False)
        if "摘录表" not in workbook.sheetnames:
            problems.append(("fatal", str(path), "-", "MISSING_SHEET"))
            workbook.close()
            continue
        ws = workbook["摘录表"]
        row_iter = ws.iter_rows()
        headers = [str(c.value).strip() if c.value is not None else ""
                   for c in next(row_iter)]
        while headers and not headers[-1]:
            headers.pop()
        if headers != conv.HEADERS[kind]:
            problems.append(("fatal", path.name, "A1", "HEADER_MISMATCH"))
            workbook.close()
            continue
        positions = {name: idx for idx, name in enumerate(headers)}
        for row_no, cells in enumerate(row_iter, start=2):
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            inspected[kind] += 1
            idcell = cells[positions["clause_no"]]
            if idcell.value is not None and idcell.data_type != "s":
                problems.append(("error", path.name, idcell.coordinate,
                                 "CLAUSE_NO_NOT_TEXT"))
            for field in (*ENUM_COLS, *LIST_COLS):
                if field not in positions:
                    continue
                cell = cells[positions[field]]
                if cell.value is None:
                    continue
                value = str(cell.value)
                if any(sep in value for sep in SUSPICIOUS_SEPARATORS):
                    # In free-text fields commas can be legitimate prose; warn,
                    # while enum columns mandate semicolon/newline separators.
                    severity = "error" if field in ENUM_COLS else "review"
                    problems.append((severity, path.name, cell.coordinate,
                                     "SUSPICIOUS_LIST_SEPARATOR_" + field))
            author = cells[positions["抽取人"]].value
            checker = cells[positions["复核人"]].value
            if checker and author and str(checker).strip() == str(author).strip():
                problems.append(("review", path.name,
                                 cells[positions["复核人"]].coordinate,
                                 "SELF_REVIEW_NEEDS_HUMAN_CONFIRMATION"))
            if kind == "rbi":
                original = cells[positions["风险等级(原始)"]].value
                unified = cells[positions["风险等级(统一)"]].value
                code = cells[positions["standard_code"]].value
                note = cells[positions["备注"]].value
                if (isinstance(code, str) and code.startswith("GB/T 26610")
                    and original and str(original).strip() not in conv.TERM_UNIFIED
                    and not (cells[positions["失效可能性等级"]].value
                             and cells[positions["失效后果等级"]].value)
                    and not note):
                    problems.append(("review", path.name,
                                     cells[positions["备注"]].coordinate,
                                     "NONSTANDARD_RISK_WORD_WITHOUT_NOTE"))
                if not code and unified:
                    problems.append(("review", path.name,
                                     cells[positions["风险等级(统一)"]].coordinate,
                                     "UNIFIED_WITHOUT_STANDARD"))
        summaries.append({"file": path.name, "data_rows": inspected[kind],
                          "header_columns": len(headers)})
        workbook.close()
    levels = Counter(x[0] for x in problems)
    print("Official XLSX nonmutating preflight:")
    for item in summaries:
        print(f"  {item['file']}: headers={item['header_columns']} "
              f"nonempty_rows={item['data_rows']}")
    print("ISSUE_COUNTS", dict(levels), "TOTAL", len(problems))
    print("ISSUE_TYPES", dict(Counter(x[3] for x in problems)))
    print("NOTE: free-text comma/dunhao marks are heuristic review tags, not proven errors.")
    for level, path, cell, issue in problems[:12]:
        print(f"  [{level}] {path} {cell}: {issue}")
    if len(problems) > 12:
        print(f"  ... {len(problems)-12} additional review tags not printed")
    print("NOTES: a self-review warning is not proof of misconduct; "
          "source text accuracy and applicability still need human review.")
    return {"fatal": levels["fatal"], "error": levels["error"],
            "review": levels["review"], "rows": dict(inspected)}


if __name__ == "__main__":
    state = check()
    raise SystemExit(1 if state["fatal"] or state["error"] else 0)
