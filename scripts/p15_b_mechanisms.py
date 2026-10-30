from __future__ import annotations
import csv, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC_XLSX = ROOT / 'upstream/02_损伤机理参考/条款摘录表_损伤机理.xlsx'
DM_PDF = ROOT / 'upstream/02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf'
ROW_MAP = ROOT / 'work/data_staging/p14/agents/A/dm30579_existing82_row_screen.csv'
SUP_COVER = ROOT / 'work/data_staging/p14/agents/C/supplement_chapter_coverage.csv'
SUP_SCAN = ROOT / 'work/data_staging/p14/agents/C/supplement_page_scan.csv'
OUT = ROOT / 'work/data_staging/p15/B_DM'
DOCS = ROOT / 'docs/data_quality'

CORRECTIONS = {
 '5.6': {
   '原文摘录': ('5.6.1 损伤描述及损伤机理：在177℃～427℃温度范围内，环烷酸对金属材料的腐蚀：2RCOOH+Fe→Fe(RCOO)₂+H₂。', '5.6.1；印刷页9；PDF页索引14；原图可视核读（机器候选，待人工逐字签核）', '在177℃～427℃温度范围内，环烷酸对金属材料的腐蚀。'),
   '损伤描述及机理': ('在177℃～427℃温度范围内，环烷酸对金属材料的腐蚀（2RCOOH+Fe→Fe(RCOO)₂+H₂）', '5.6.1；印刷页9；PDF页索引14；原图可视核读（机器候选，待人工逐字签核）', '在177℃～427℃温度范围内，环烷酸对金属材料的腐蚀。'),
   '温度下限': (177, '5.6.1；印刷页9；PDF页索引14；原图可视核读（机器候选，待人工逐字签核）', '177℃'),
   '温度上限': (427, '5.6.1；印刷页9；PDF页索引14；原图可视核读（机器候选，待人工逐字签核）', '427℃'),
   '其他适用条件': ('通用机理温区：GB/T 30579-2022 §5.6.1为177℃～427℃；§5.6.4(b)的218℃～400℃仅指轻相介质使用温度下腐蚀较常见，不替代通用机理温区。腐蚀速率还受酸值、硫含量、流速、相态和材料影响；§5.6.4(f)称钼含量下限为2.5%（质量分数）时可提高耐蚀性。', '5.6.1、5.6.4(b)；印刷页9；PDF页索引14；原图可视核读（机器候选，待人工逐字签核）', '轻相介质218℃～400℃范围内腐蚀较为常见。'),
 },
 '5.23': {
   '温度下限': (None, '5.23.4(e)、5.23.5(a)；印刷页33–34；PDF页索引38–39；页图机器核读（待人工核定）', '49℃～66℃是局部析出/沉积现象温度窗口，不是整体机理温度范围。'),
   '温度上限': (None, '5.23.4(e)、5.23.5(a)；印刷页33–34；PDF页索引38–39；页图机器核读（待人工核定）', '49℃～66℃是局部析出/沉积现象温度窗口，不是整体机理温度范围。'),
   '其他适用条件': ('保留对象范围：介质为含硫氢化铵（NH₄HS）的酸性水；NH₄HS浓度超过2%（质量分数）及流态/流向变化等是形态与严重程度条件。49℃～66℃仅为特定加氢反应流出物等场景中气相NH₄HS析出/沉积的局部现象窗口，不能作为全机理适用温区。§5.23.3将碳钢描述为耐蚀性弱，并称300系列、双相、铝及镍基合金耐蚀性较强；材料列举不等于所有列材质均高敏感。', '5.23.2–5.23.5；印刷页32–34；PDF页索引37–39；原图机器可视核读（待人工核定）', '硫氢化铵浓度超过2%（质量分数）。'),
 },
 '6.7': {
   '温度上限': (None, '6.7.4；印刷页52；PDF页索引57；原图机器核读（待人工核定）', 'ambient to 150℃ applies to HBB/HIC/SOHIC subtype; SSC below 82℃; no common global range.'),
   '其他适用条件': ('湿H₂S环境下需按子类型分别判断：氢鼓泡/HIC/SOHIC通常涉及环境温度至150℃、有时更高；SSC通常在82℃以下。0.0003 MPa表示H₂S气相分压门槛，不是设备操作总压力；S/T总压力字段保持空值。原材料易感性与H₂S分压、pH及硬度/应力等条件关联。', '6.7.3–6.7.4；印刷页51–52；PDF页索引56–57；原图机器可视核读（待人工核定）', '0.0003 MPa is H₂S partial pressure.'),
 }
}
FIELD_META = {
 'A':'REQUIRED','B':'REQUIRED','C':'REQUIRED','D':'REQUIRED','E':'REQUIRED','F':'REQUIRED',
 'G':'REQUIRED','H':'CONDITIONAL_SUPPORTED','I':'CONDITIONAL_SUPPORTED','J':'CONDITIONAL_SUPPORTED',
 'K':'CONDITIONAL_SUPPORTED','L':'CONDITIONAL_SUPPORTED','M':'CONDITIONAL_SUPPORTED','N':'CONDITIONAL_SUPPORTED',
 'O':'CONDITIONAL_SUPPORTED','P':'CONDITIONAL_SUPPORTED','Q':'CONDITIONAL_SUPPORTED','R':'CONDITIONAL_SUPPORTED',
 'S':'SOURCE_UNSPECIFIED','T':'SOURCE_UNSPECIFIED','U':'CONDITIONAL_SUPPORTED','V':'CONDITIONAL_SUPPORTED',
 'W':'REVIEW_NEEDED','X':'REVIEW_NEEDED','Y':'CONDITIONAL_SUPPORTED'
}


def sha(path):
 h=hashlib.sha256()
 with path.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
 return h.hexdigest()

def csv_write(path, rows, fields):
 path.parent.mkdir(parents=True,exist_ok=True)
 with path.open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore'); w.writeheader(); w.writerows(rows)

def load_rows(path):
 with path.open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))

def row_to_excel_id(r): return str(r.get('clause_id','')).replace('GBT30579-2022-','')

# Validate source inventory and derive all 656 G-N review slots.
wb=load_workbook(SRC_XLSX)
ws=wb['摘录表']
headers=[c.value for c in ws[1]]
assert len(headers)==25 and headers[0]=='clause_id'
rows={str(ws.cell(i,1).value).replace('GBT30579-2022-',''):i for i in range(2,ws.max_row+1) if ws.cell(i,1).value}
assert len(rows)==82, len(rows)
locator={row_to_excel_id(r):r for r in load_rows(ROW_MAP)}
assert set(rows)==set(locator), (set(rows)-set(locator),set(locator)-set(rows))
# produce candidate from original and propose only source-backed corrections.
candidate=load_workbook(SRC_XLSX)
cws=candidate['摘录表']
changelog=[]
for cid, changes in CORRECTIONS.items():
 r=rows[cid]
 for field,(new,source,evidence) in changes.items():
  col=headers.index(field)+1; old=cws.cell(r,col).value
  if old == new: continue
  cws.cell(r,col).value=new
  changelog.append({'clause_id':cid,'excel_row':r,'field':field,'old_value':'' if old is None else old,'new_value':'' if new is None else new,
   'action':'PROPOSED_CORRECTION','standard_code':'GB/T 30579-2022','subclause':source.split('；')[0], 'printed_page':source.split('；')[1].replace('印刷页',''),
   'pdf_page_index':source.split('；')[2].replace('PDF页索引',''), 'evidence_locator':source,'evidence_quote':evidence,'method':'本地渲染原页图像+机器视觉核读；不等于人工签核',
   'reviewer_status':'REVIEW_NEEDED','human_reviewer':'','reason':'局部条件与整个机理边界分离；待标准原图人工逐字复核后方可入正式库'})
# Ensure no human reviewer was forged; preserve empty X.
for r in range(2,cws.max_row+1): cws.cell(r,24).value=None
OUT.mkdir(parents=True,exist_ok=True)
xlsx_out=OUT/'条款摘录表_损伤机理_P15_B_candidate.xlsx'
candidate.save(xlsx_out)

# OCR locator evidence is attached per slot but never treated as verified ground truth.
ocr_path=OUT/'DM_SUBCLAUSE_OCR_LOCATORS.jsonl'
ocr_map={}
if ocr_path.exists():
 for line in ocr_path.read_text(encoding='utf-8').splitlines():
  x=json.loads(line); ocr_map[(x['clause_id'],x['field_letter'])]=x
# 656-slot per-field ledger; all are explicitly review-needed unless a proposed correction is scoped to a specific cell.
field_names={chr(ord('G')+i):headers[6+i] for i in range(8)}
triage=[]
condition_fields={'O':'适用材质','P':'适用介质','Q':'温度下限','R':'温度上限','S':'压力下限','T':'压力上限','U':'适用设备类型','V':'其他适用条件'}
condition_subclause={'O':'.3','P':'.1/.4','Q':'.1/.4','R':'.1/.4','S':'.4','T':'.4','U':'.5','V':'.4/.5'}
for cid,r in sorted(rows.items(),key=lambda x:(int(x[0].split('.')[0]),int(x[0].split('.')[1]))):
 m=locator[cid]; printed=m.get('printed_page','');
 try: pdf_index=int(m['pdf_page_index']); pdf_page=pdf_index+1
 except: pdf_index=''; pdf_page=''
 for col,field in field_names.items():
  colno=headers.index(field)+1; old=ws.cell(r,colno).value; new=cws.cell(r,colno).value
  proposed=(old!=new)
  triage.append({'clause_id':cid,'excel_row':r,'field_letter':col,'field_name':field,'field_class':FIELD_META[col],
    'old_value':'' if old is None else old,'candidate_value':'' if new is None else new,
    'source_standard':'GB/T 30579-2022','source_subclause':f'{cid}.{ord(col)-ord("G")+1}',
    'printed_page_locator':printed,'pdf_page_index_0based':pdf_index,'pdf_page_number_1based':pdf_page,
    'source_bbox_or_table_cell':'NOT_CAPTURED','evidence_quote':'','ocr_heading_found':ocr_map.get((cid,col),{}).get('ocr_heading_found',False),'ocr_status':ocr_map.get((cid,col),{}).get('ocr_status','OCR_LOCATOR_NOT_RUN'),'ocr_snippet':ocr_map.get((cid,col),{}).get('ocr_snippet',''),'numeric_quantity_semantics':'详见该机理子条款；需逐值确认对象/单位/边界',
    'enum_mapping_status':'需要按冻结enums.json逐格复核' if col in ('I','O','P','U') else 'N/A',
    'evidence_status':'PROPOSED_CORRECTION' if proposed else 'REVIEW_NEEDED',
    'reviewer_status':'NOT_HUMAN_VERIFIED','reason':'仅完成章节起始页/八段子条号机器映射；未完成该字段逐格原文语义比对' if not proposed else '候选变更见P15_B_FIELD_CHANGELOG.csv；必须人工核图确认'})
 # Carry the independent applicability/operating-condition fields in the same cell-level ledger.
 for col,field in condition_fields.items():
  colno=headers.index(field)+1; old=ws.cell(r,colno).value; new=cws.cell(r,colno).value
  suffix=condition_subclause[col]
  triage.append({'clause_id':cid,'excel_row':r,'field_letter':col,'field_name':field,'field_class':FIELD_META[col],
    'old_value':'' if old is None else old,'candidate_value':'' if new is None else new,
    'source_standard':'GB/T 30579-2022','source_subclause':cid+suffix,
    'printed_page_locator':printed,'pdf_page_index_0based':pdf_index,'pdf_page_number_1based':pdf_page,
    'source_bbox_or_table_cell':'NOT_CAPTURED','evidence_quote':'','ocr_heading_found':'','ocr_status':'OCR_SUBSECTION_FAMILY_LOCATOR_ONLY','ocr_snippet':'',
    'numeric_quantity_semantics':{'Q':'需确认是否为机理温度边界而非局部现象温度','R':'需确认是否为机理温度边界而非局部现象温度','S':'需确认压力类型及介质/气相对象','T':'需确认压力类型及介质/气相对象'}.get(col,'需确认条款对象与适用性，不以出现词语等同适用'),
    'enum_mapping_status':'需要按冻结enums.json逐项复核' if col in ('O','P','U') else 'N/A',
    'evidence_status':'PROPOSED_CORRECTION' if old!=new else 'REVIEW_NEEDED','reviewer_status':'NOT_HUMAN_VERIFIED',
    'reason':'适用条件专项审计未逐值完成；定位至机理属性章节族，需逐条人工核对源页' if old==new else '候选变更见P15_B_FIELD_CHANGELOG.csv；必须人工核图确认'})
csv_write(OUT/'DM_FIELD_TRIAGE.csv',triage,list(triage[0]))

# Supplement chapter inventory: expand ranges from P14 coverage; classification is not ingestion authorization.
cover=load_rows(SUP_COVER)
scan=load_rows(SUP_SCAN)
scan_counts={}
for r in scan: scan_counts[r['standard_code']]=scan_counts.get(r['standard_code'],0)+1
specs={
 'GB/T 150.1-2024':('GBT150.1-2024_压力容器第1部分通用要求.pdf',51,'auxiliary_background','压力容器通用建造/材料/设计语境；第4章仅通用失效模式类别，未形成30579式服役损伤机理记录。'),
 'GB/T 150.3-2024':('GBT150.3-2024_压力容器第3部分设计.pdf',287,'auxiliary_background','结构设计、强度/稳定性和低温脆断校核是设计约束背景，不得推导服役腐蚀风险或检验周期。'),
 'GB/T 150.4-2024':('GBT150.4-2024_压力容器第4部分制造检验和验收.pdf',67,'auxiliary_background','制造、检验和验收内容可作制造质量辅助索引；制造NDT不等于在役损伤检测监测方法。'),
 'GB/T 151-2026':('GBT151-2026_热交换器.pdf',339,'auxiliary_background','热交换器设计/制造与振动等可能支撑设备背景；未有经核对的条款证明可直接映射DM八字段。'),
 'GB/T 21447-2018':('GBT21447-2018_钢质管道外腐蚀控制规范.pdf',17,'auxiliary_background','管道外腐蚀控制相关；4.3环境分类及5/6防腐控制仅保留条件性辅助索引，不是新增30579机理。')}
expanded=[]
for c in cover:
 code=c['standard_code']; label=c['chapter_or_annex']; title=c['chapter_title_or_status']; file,pages,base,rationale=specs[code]
 nums=[]
 if re.fullmatch(r'\d+-\d+',label):
  a,b=map(int,label.split('-')); nums=[str(i) for i in range(a,b+1)]
 elif label in ('A-F','A-G','A-E','A-S'):
  a,b=label.split('-'); nums=[chr(i) for i in range(ord(a),ord(b)+1)]
 else: nums=[label]
 for n in nums:
  exact=(n==label)
  t=title if exact else '章名待逐页目录视觉核定'
  if code=='GB/T 21447-2018' and n=='4': t='基本规定（P14索引指向4.3及表1–4，单元格未逐格核验）'
  if code=='GB/T 21447-2018' and n=='5': t='防腐层设计（辅助控制背景）'
  if code=='GB/T 21447-2018' and n=='6': t='阴极保护设计（辅助控制背景）'
  if code=='GB/T 151-2026' and n=='4': t='失效模式（目录片段；范围待视觉核定）'
  if code=='GB/T 151-2026' and n=='A' : t='附录A；章名待视觉核定'
  if code=='GB/T 151-2026' and n=='C' : t='附录C：流体诱发振动（P14标题线索，待核）'
  if n=='4' and code=='GB/T 21447-2018': disposition='auxiliary_background'; why='外腐蚀环境分级可作环境/控制检索索引；表1–4行列值未视觉逐格核验，不形成DM候选。'
  elif n in ('5','6') and code=='GB/T 21447-2018': disposition='auxiliary_background'; why='防腐层/阴极保护设计属于预防控制背景；不等同于损伤机理条目。'
  elif code=='GB/T 151-2026' and n=='C': disposition='auxiliary_background'; why='流体诱发振动可作设备失效背景；未证明其可直接映射GB/T30579损伤机理字段。'
  elif n=='4' and code=='GB/T 150.1-2024': disposition='out_of_scope'; why='通用失效模式类别无具体服役腐蚀机理和工况边界。'
  elif n in ('7','8','9','10') and code=='GB/T 21447-2018': disposition='REVIEW_NEEDED'; why='目录标题文本损坏/未视觉核定，先保留待核；未形成正式或候选条目。'
  elif '待' in t or (not exact and c.get('coverage_status','').startswith('MACHINE')): disposition='REVIEW_NEEDED'; why='P14只有章节范围/机器目录扫描，具体章名及页边界尚未视觉复核；不得猜写章名或提取条款。'
  else: disposition=base; why=rationale
  expanded.append({'standard_code':code,'source_file':file,'source_sha256':sha(ROOT/'upstream/02_损伤机理参考/补充标准'/file),'chapter_or_annex':n,'chapter_title_or_status':t,'pdf_pages_in_standard':pages,
    'printed_page_start':'','printed_page_end':'','pdf_page_index_start_0based':'','pdf_page_index_end_0based':'',
    'disposition':disposition,'direct_dm_candidate':'NO','auxiliary_index':'YES' if disposition=='auxiliary_background' else 'NO',
    'source_quality':'P14全页机器扫描；原生文本受CMap/乱码影响；本轮未完成逐章人工视觉复核',
    'rationale':why,'human_review_status':'NOT_DONE','reviewer':''})
# ensure full counts by code and no empty dispositions
assert all(x['disposition'] and x['rationale'] for x in expanded)
csv_write(DOCS/'P15_B_SUPPLEMENT_CHAPTER_DISPOSITION.csv',expanded,list(expanded[0]))
# Auxiliary knowledge index only: no unsupported verbatim extraction or DM clause records.
aux=[]
for e in expanded:
 if e['disposition']=='auxiliary_background' or (e['standard_code']=='GB/T 21447-2018' and e['chapter_or_annex'] in ('4','5','6')):
  aux.append({'record_id':f"P15-B-AUX-{e['standard_code'].replace('/','').replace(' ','-')}-{e['chapter_or_annex']}",
   'record_type':'auxiliary_knowledge_index','standard_code':e['standard_code'],'source_file':e['source_file'],
   'chapter_or_annex':e['chapter_or_annex'],'title':e['chapter_title_or_status'],'topic':e['rationale'],
   'source_locator_status':'CHAPTER_ONLY_NO_VERIFIED_PAGE_RANGE','raw_quote':'','bbox_or_table_cells':None,
   'disposition':'auxiliary_background','review_status':'REVIEW_NEEDED','human_reviewer':'',
   'not_for':'DM正式表/正式clauses.jsonl；不能推导新机理、风险等级或检验周期'})
(OUT/'supplemental_index.jsonl').write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in aux),encoding='utf-8')

csv_write(DOCS/'P15_B_FIELD_CHANGELOG.csv',changelog,list(changelog[0]))
# hashes and tallies
status={}
for r in triage: status[r['evidence_status']]=status.get(r['evidence_status'],0)+1
std_counts={}
for e in expanded: std_counts[e['standard_code']]=std_counts.get(e['standard_code'],0)+1
summary={'source_xlsx_sha256':sha(SRC_XLSX),'source_pdf_sha256':sha(DM_PDF),'candidate_xlsx_sha256':sha(xlsx_out),
 'dm_rows':len(rows),'field_ledger_rows':len(triage),'g_n_field_slots':sum(1 for r in triage if r['field_letter'] in field_names),'condition_field_slots':sum(1 for r in triage if r['field_letter'] in condition_fields),'ocr_headings_found':sum(1 for r in triage if r['field_letter'] in field_names and r['ocr_heading_found']),'ocr_headings_missing':sum(1 for r in triage if r['field_letter'] in field_names and not r['ocr_heading_found']),'field_slots_by_status':status,'proposed_cell_changes':len(changelog),
 'candidate_clause_ids_preserved':len(rows),'reviewer_cells_blank':sum(1 for r in rows.values() if not cws.cell(r,24).value),
 'supplement_page_total':sum(v[1] for v in specs.values()),'supplement_chapter_rows':len(expanded),'supplement_chapters_by_standard':std_counts,'supplement_sources':{code:{'file':specs[code][0],'pages':specs[code][1],'sha256':sha(ROOT/'upstream/02_损伤机理参考/补充标准'/specs[code][0])} for code in specs},
 'direct_supplement_dm_candidates':0,'human_review_completed':False,'status':'REVIEW_BLOCKED'}
(OUT/'run_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')

# report includes machine/human distinction and explicit limitations.
report=f'''# P15-B 损伤机理字段与补充标准质量报告\n\n状态：**REVIEW_BLOCKED**（候选输出已整理；不能发布为人工复核完成数据）\n\n生成时间：2026-10-29（本机时区）；生成脚本：`scripts/p15_b_mechanisms.py`。\n\n## 输入与冻结\n\n- DM上游Excel：`upstream/02_损伤机理参考/条款摘录表_损伤机理.xlsx`，SHA-256 `{summary['source_xlsx_sha256']}`。\n- GB/T 30579-2022 PDF：SHA-256 `{summary['source_pdf_sha256']}`。上游PDF/Excel均只读。\n- 候选Excel：`work/data_staging/p15/B_DM/条款摘录表_损伤机理_P15_B_candidate.xlsx`，SHA-256 `{summary['candidate_xlsx_sha256']}`。\n- 82/82模式编号原样保留；工作簿原4 sheet/25列保留。复核人X列全空；抽取人W“代航”原样保留但不认定身份有效。\n\n## 全量字段映射范围\n\n- G–N：82×8 = **656**个八段式字段；O–V中的材质/介质/温压/设备/其他适用条件再登记656个cell。`DM_FIELD_TRIAGE.csv`共1312条cell记录，覆盖上述两组而不混为同一验收计数。\n- {status.get('REVIEW_NEEDED',0)}个字段槽保持`REVIEW_NEEDED`；{status.get('PROPOSED_CORRECTION',0)}个字段槽对应候选更正。163页本地渲染图已用Tesseract 5.5.2 `chi_sim+eng/psm6`批处理，656个槽中{summary['ocr_headings_found']}个找到机器OCR子条标题，{summary['ocr_headings_missing']}个未找到。OCR定位片段见`DM_SUBCLAUSE_OCR_LOCATORS.jsonl`，完整页OCR见`DM_PAGE_OCR.jsonl`。当前**没有**为字段伪造独立bbox、表格单元或可信逐字段引文；OCR片段受错识/跨栏影响，只作导航，不是证据签字。来源页定位不是逐字段核验。其余G–N及A–F/O–V具体适用性都需原图逐格核对。\n- 本机图像读取只作为机器辅助，所有变更状态`REVIEW_NEEDED`，不存在`HUMAN_VERIFIED`。\n\n## 已提出的实质变更\n\n候选共{len(changelog)}个单元格变更，详见`docs/data_quality/P15_B_FIELD_CHANGELOG.csv`。\n\n1. **5.6 环烷酸腐蚀**：原F/G及Q/R误写177–242℃。§5.6.1图像显示机理温区177–427℃；§5.6.4(b)的218–400℃限定“轻相介质使用温度下腐蚀常见”，不是整个机制温区。候选将F/G/Q/R改为177–427℃，V把两个条件对象分开。来源：印刷页9，PDF 0-based index 14（图`upstream/02_损伤机理参考/_rendered/single_014.png`）。\n2. **5.23 酸性水腐蚀（碱式酸性水）**：将Q/R的49–66℃清空；V标明这是特定析出/沉积局部窗口，不是整个机理过滤区间。保留NH₄HS浓度/流态条件；“耐蚀材料”与“高易感材料”语义分开。来源定位印刷页32–34，PDF indexes 37–39，候选证据限机器图像核读，须逐页人工核。\n3. **6.7 湿H₂S破坏**：R=150℃清空；V将子类型温度适用性分开，并指出0.0003 MPa为H₂S气相分压，不是总操作压力。S/T继续空。需人工核清不同子类型和单位边界后再用。\n\n上述更改是**带来源的机器候选**，不是经人工签核的事实。F逐字摘录门禁同样未关闭。\n\n## 五份补充标准章节处理\n\n- 原PDF合计761页；按P14索引扩展为{len(expanded)}章/附录单元，均有分类/理由；逐标准计数：{json.dumps(std_counts,ensure_ascii=False)}。各PDF SHA-256记录于`P15_B_SUPPLEMENT_CHAPTER_DISPOSITION.csv`及`work/data_staging/p15/B_DM/run_summary.json`。\n- 分类计数：{json.dumps({k:sum(1 for r in expanded if r['disposition']==k) for k in sorted(set(r['disposition'] for r in expanded))},ensure_ascii=False)}。未知章名、标题损坏或页界未视觉核实者保留`REVIEW_NEEDED`，页码范围不臆填。\n- 150.1/150.3/150.4/151作为设计、制造、设备/振动等辅助背景索引；不会因标准中出现“失效模式”“检测”“防腐”就生成DM条目。\n- 21447 §4环境分类、§5防腐层、§6阴极保护仅辅助索引；4.3表1–4未逐格核表，页码定位不能替代格值，故**0**条补充标准DM候选进入候选Excel/正式JSONL。`supplemental_index.jsonl`不含未经核对的引文。\n\n## 质量门与未完成项\n\n- 保留82模式、候选结构校验：通过（由本脚本assert及专项测试复核）。\n- 656/656字段“有子条号+页定位”；独立字段原文/数值/枚举适用性证据：**未完成**；REVIEW_NEEDED字段不得计作来源有据字段。\n- 5.6、5.23、6.7局部温度/分压语义：候选已处理，待人工复核PDF原图；没有编辑6.7 S/T。5.23 O材料枚举不删改，不把耐蚀材质解释成高易感材质。\n- 五份PDF按章分类/理由完整；章标题/页边界的全部原文级复核未完成，21447表格逐格复核未完成。\n- 无真人reviewer、无真实签字、无人工完成率声明；状态保持`REVIEW_BLOCKED`。\n\n## 可复现命令\n\n```powershell\npython scripts/p15_b_mechanisms.py\npython -m pytest tests/test_p15_b_mechanisms.py -q\n```\n'''
(DOCS/'P15_B_MECHANISM_FIELD_REVIEW.md').write_text(report,encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))

