from __future__ import annotations
import csv, hashlib, json, re, shutil
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'upstream/01_RBI标准/条款摘录表_RBI.xlsx'
PDF = ROOT / 'upstream/01_RBI标准/GBT26610.3-2014_风险的定性分析方法.pdf'
OUT = ROOT / 'work/data_staging/p15/A_RBI'
OUT.mkdir(parents=True, exist_ok=True)
CANDIDATE = OUT / '条款摘录表_RBI_P15_A_candidate.xlsx'
TRIAGE = OUT / 'RBI_FIELD_TRIAGE.csv'
CHANGELOG = OUT / 'P15_A_FIELD_CHANGELOG.csv'
SOURCE_LEDGER = OUT / 'P15_A_SOURCE_LEDGER.csv'
UNRESOLVED = OUT / 'P15_A_UNRESOLVED.csv'
MATRIX = OUT / 'matrix_266103_pdfindex13.png'


def sha(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def s(v): return '' if v is None else str(v).strip()
def norm(v): return re.sub(r'\\s+','',s(v))

def load_rows(path):
    wb=load_workbook(path, read_only=True, data_only=True)
    ws=wb['摘录表']; headers=[c.value for c in ws[1]]
    data=[]
    for excel_row,row in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
        vals=list(row[:len(headers)])
        if vals and s(vals[0]): data.append((excel_row,dict(zip(headers,vals))))
    return headers,data

headers, records=load_rows(SOURCE)
assert len(headers)==23 and len(records)==485, (len(headers),len(records))
shutil.copy2(SOURCE,CANDIDATE)
wb=load_workbook(CANDIDATE)
ws=wb['摘录表']
# Evidence: standard's Figure 1 visual on 0-based physical PDF page 13.
# Colors read from the rendered source figure, not inferred from risk rules.
colors={
 5:['中高风险','中高风险','中高风险','高风险','高风险'],
 4:['中风险','中风险','中高风险','中高风险','高风险'],
 3:['低风险','低风险','中风险','中高风险','高风险'],
 2:['低风险','低风险','中风险','中风险','中高风险'],
 1:['低风险','低风险','中风险','中风险','中高风险'],
}
letters='ABCDE'
expected={(str(l)+c):colors[l][i] for l in range(1,6) for i,c in enumerate(letters)}
# Verify existing matrix cell assertions against source figure before any candidate edit.
mat=[(r,d) for r,d in records if s(d.get('standard_code'))=='GB/T 26610.3-2014' and s(d.get('clause_no'))=='7']
assert len(mat)==25, f'expected 25 risk matrix rows, got {len(mat)}'
for excel_row,d in mat:
    key=s(d['失效可能性等级'])+s(d['失效后果等级'])
    assert key in expected, (excel_row,key)
    assert s(d['风险等级(原始)'])==expected[key], (excel_row,key,d['风险等级(原始)'],expected[key])
    unified={'低风险':'L1','中风险':'L2','中高风险':'L3','高风险':'L4'}[expected[key]]
    assert s(d['风险等级(统一)'])==unified, (excel_row,key,d['风险等级(统一)'],unified)

# Render the source PDF figure for this reproducible evidence package.
import fitz
pdf=fitz.open(PDF)
assert len(pdf)>13 and '图1' in pdf[13].get_text().replace(' ','')
pdf[13].get_pixmap(matrix=fitz.Matrix(2,2), alpha=False).save(MATRIX)
pdf_pages=pdf[13].get_text()

triage=[]; changelog=[]; source_ledger=[]; unresolved=[]
field_meta={
 'clause_id':('REQUIRED','主键必需；内容仍待原文逐字/身份核验'),
 'standard_code':('REQUIRED','标准标识必需；来源文件可定位'),
 'clause_no':('REQUIRED','条款标识必需；逐条定位未完成人审'),
 'page':('REQUIRED','印刷页码必需；与PDF物理页不同，需逐条核对'),
 'title':('REQUIRED','标题必需；应与来源标题一致'),
 '原文摘录':('REVIEW_NEEDED','原文必须逐字且<=300字；表格/乱码行待逐单元格核对'),
 '失效可能性等级':('CONDITIONAL_SUPPORTED','仅标准确有分级时适用；.4概率结果不可映射为此等级'),
 '失效后果等级':('CONDITIONAL_SUPPORTED','仅标准确有后果等级时适用'),
 '风险等级(原始)':('CONDITIONAL_SUPPORTED','只对标准提供或.3矩阵映射行适用'),
 '风险等级(统一)':('CONDITIONAL_SUPPORTED','应由标准分级或映射表导出；.3矩阵可机器校验'),
 '检验周期(年)':('CONDITIONAL_SUPPORTED','只在原条款定义可结构化固定/条件周期时填写；动态RBI需保留空'),
 '推荐检验方法':('CONDITIONAL_SUPPORTED','来源明确方法时才填写；其余待核方法语义'),
 '适用材质':('CONDITIONAL_SUPPORTED','仅来源明确指定材质/适用对象时填写；不以推断填补'),
 '适用介质':('CONDITIONAL_SUPPORTED','仅来源明确指定介质时填写'),
 '温度下限':('CONDITIONAL_SUPPORTED','仅来源明确给出与本字段语义一致的操作温度下限时填写'),
 '温度上限':('CONDITIONAL_SUPPORTED','仅来源明确给出与本字段语义一致的操作温度上限时填写'),
 '压力下限':('CONDITIONAL_SUPPORTED','仅来源明确给出与本字段语义一致的操作压力下限时填写'),
 '压力上限':('CONDITIONAL_SUPPORTED','仅来源明确给出与本字段语义一致的操作压力上限时填写'),
 '适用设备类型':('CONDITIONAL_SUPPORTED','来源明确限定设备类型时填写'),
 '其他适用条件':('CONDITIONAL_SUPPORTED','仅设备/工况判据；编辑说明/表格拆行说明应归备注'),
 '抽取人':('REVIEW_NEEDED','保留既有身份字符串；不能由机器确认真实责任人/学号格式'),
 '复核人':('REVIEW_NEEDED','未签不得代签；空值是人审门禁'),
 '备注':('CONDITIONAL_SUPPORTED','来源定位/限制说明可以记录在此；人工审核状态不得伪造'),
}

cache=ROOT/'work/data_staging/p14/page_cache/native_pages.jsonl'
cache_rows=[]
if cache.exists():
  import json as _json
  with cache.open(encoding='utf-8') as _f:
    for _line in _f:
      try: cache_rows.append(_json.loads(_line))
      except Exception: pass
for excel_row,d in records:
    code=s(d.get('standard_code')); clause=s(d.get('clause_no')); cid=s(d.get('clause_id'))
    is_draft=code=='GB/T 26610.1-2022'
    for field in headers:
        value=d.get(field)
        category,base_reason=field_meta[field]
        reason=base_reason
        if field in ('适用材质','适用介质','温度下限','温度上限','压力下限','压力上限','适用设备类型','其他适用条件','检验周期(年)') and (value is None or s(value)==''):
            category='SOURCE_UNSPECIFIED'
            reason='当前单元格为空；本轮尚未定位到能够支持该字段的条款原文，保持空值。'
        elif field=='复核人' and not s(value):
            category='REVIEW_NEEDED'; reason='空白是待真人交叉复核状态；不得机器代签。'
        elif field=='原文摘录' and len(s(value))>300:
            reason='超过300字限制；机器标记待按来源条款/表格单元拆分，未自动截断。'
        elif field=='其他适用条件' and s(value) and (('拆行' in s(value) or '逐格拆行' in s(value) or '对应 -' in s(value)) or excel_row in (450,451,454,456)):
            category='PROPOSED_CORRECTION'
            reason=('TSG 21 条款范围/条件存在针对性遗漏或错位，候选修正已由原文定位；保留人工审签门禁。' if excel_row in (450,451,454,456) else '图1矩阵拆行映射/录入说明并非设备工况适用条件；已迁至备注并保留原说明。')
        elif field=='原文摘录' and excel_row in range(144,154):
            category='PROPOSED_CORRECTION'; reason='GB/T 26610.5表2/表3行摘录已有逐行渲染图交叉核查候选；仍需真人签核，变更留在隔离副本。'
        elif field=='原文摘录' and re.search(r'表\\s*\\d|表\\d|对应.*等级|对应.*因子|图1',s(value)):
            reason='表格单元/解释性转述模式；属于启发式待核，不等同于已判错。'
        triage.append({'excel_row':excel_row,'clause_id':cid,'standard_code':code,'clause_no':clause,'field':field,'value_state':'NONEMPTY' if value is not None and s(value) else 'EMPTY','field_class':category,'reason':reason,'draft_status':'DRAFT_UNSIGNED' if is_draft else 'EXISTING_UNREVIEWED','source_pdf':'','printed_page':d.get('page') or '', 'pdf_page_index':'','table_id':'','table_row':'','table_column':'','evidence_quote':'','machine_review':'NOT_VERIFIED'})

# Move only explicit Figure 1 splitting/editorial notes (25 matrix cells) from T to W.
for excel_row,d in mat:
    old_t=s(d.get('其他适用条件')); old_w=s(d.get('备注'))
    assert '风险矩阵共25格' in old_t and '本行为' in old_t and '对应 -' in old_t, (excel_row,old_t)
    new_t=None
    tag='[P15 A语义归位] 原T列为图1矩阵拆行/后缀映射说明，不是设备适用条件；保留原说明以便追溯：'+old_t
    corrected_w=old_w.replace('图1在正文第12页','图1在正文第14页（PDF index 13）')
    new_w=(corrected_w+'；' if corrected_w else '')+tag
    # T column 20; W column 23 in the 23-column source template.
    ws.cell(excel_row,20).value=new_t
    ws.cell(excel_row,23).value=new_w
    loc={'excel_row':excel_row,'clause_id':s(d['clause_id']),'standard_code':s(d['standard_code']),'clause_no':'7','pdf_file':'upstream/01_RBI标准/GBT26610.3-2014_风险的定性分析方法.pdf','printed_page':14,'pdf_page_index':13,'bbox':'Figure 1 image; approx x=168..300pt,y=98..240pt on 595x842pt page; table cells raster/vector figure (not native text blocks)','table_id':'图1 风险矩阵','table_row':s(d['失效可能性等级']),'table_column':s(d['失效后果等级']),'element_name':'五级失效可能性 × A-E后果的颜色分区','evidence_quote':'图1 风险矩阵；标准第7章正文称组合失效可能性等级与失效后果等级后在图1确定风险级别。视觉读取矩阵颜色：'+s(d['失效可能性等级'])+s(d['失效后果等级'])+'='+expected[s(d['失效可能性等级'])+s(d['失效后果等级'])],'evidence_level':'E2_VISUAL_SOURCE_PAGE_MACHINE_RENDER','review_required':'TRUE_HUMAN_REQUIRED','source_pdf_sha256':sha(PDF)}
    for field,old,new in [('其他适用条件',old_t,''),('备注',old_w,new_w)]:
        changelog.append({'excel_row':excel_row,'clause_id':s(d['clause_id']),'field':field,'old_value':old,'new_value':new,'action':'MOVE_EDITORIAL_NOTE_TO_REMARKS' if field=='其他适用条件' else 'CORRECT_PAGE_LOCATOR_AND_PRESERVE_TRACE','why':'T列的描述仅说明图1风险矩阵拆行索引与ID后缀，不描述设备或工况适用条件；移入备注。同时将旧备注中的图1印刷页12更正为源图实际印刷页14/PDF index13。','pdf_file':loc['pdf_file'],'printed_page':14,'pdf_page_index':13,'bbox':loc['bbox'],'table_id':'图1 风险矩阵','table_row':loc['table_row'],'table_column':loc['table_column'],'element_name':loc['element_name'],'evidence_quote':loc['evidence_quote'],'evidence_level':loc['evidence_level'],'review_required':'TRUE_HUMAN_REQUIRED','source_pdf_sha256':loc['source_pdf_sha256'],'source_xlsx_row':excel_row,'source_xlsx_sha256':sha(SOURCE)})
    source_ledger.append(loc)

# Import four row-specific TSG scope fixes from the isolated source-evidence workstream.
# This CSV contains exact clauses/pages and is still explicitly marked human-review-required.
tsg_findings=OUT/'research/266102_tsg/field_findings.csv'
if tsg_findings.exists():
  with tsg_findings.open(encoding='utf-8-sig',newline='') as _f:
    for item in csv.DictReader(_f):
      if item.get('action')!='PROPOSED_CORRECTION' or item.get('field','').startswith('其他适用条件') is False:
        continue
      rowno=int(item['excel_row']); old=s(item['old_value']); new=s(item['new_value'])
      # The research row must still match the frozen input; fail closed if it drifted.
      assert s(ws.cell(rowno,20).value)==old, (rowno,ws.cell(rowno,20).value,old)
      ws.cell(rowno,20).value=new
      pdf_file='upstream/01_RBI标准/补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf'
      source_ledger.append({'excel_row':rowno,'clause_id':item['clause_id'],'standard_code':'TSG 21-2016','clause_no':item['clause_id'].split('TSG21-2016-',1)[-1],'printed_page':item['printed_page'],'pdf_file':pdf_file,'field':'其他适用条件','quote':item['evidence_quote'],'quote_chars':len(item['evidence_quote']),'quote_mode':'CLAUSE_SCOPE_NATIVE_TEXT_MACHINE_LOCATED','evidence_level':item['evidence_level'],'quote_exactness':'SOURCE_QUOTE_MACHINE_CAPTURE_NOT_HUMAN_SIGNED','human_review':'REQUIRED','pdf_page_index':item['pdf_page_index'],'bbox':item['bbox'],'table_id':item['table_id'],'table_row':item['table_row'],'table_column':item['table_column'],'element_name':item['element_name'],'evidence_quote':item['evidence_quote'],'review_required':item['review_required'],'source_pdf_sha256':item['source_sha256']})
      changelog.append({'excel_row':rowno,'clause_id':item['clause_id'],'field':'其他适用条件','old_value':old,'new_value':new,'action':'UPDATE_SCOPE_CONDITION_FROM_SOURCE','why':item['why'],'pdf_file':pdf_file,'printed_page':item['printed_page'],'pdf_page_index':item['pdf_page_index'],'bbox':item['bbox'],'table_id':item['table_id'],'table_row':item['table_row'],'table_column':item['table_column'],'element_name':item['element_name'],'evidence_quote':item['evidence_quote'],'evidence_level':item['evidence_level'],'review_required':'TRUE_HUMAN_REQUIRED','source_pdf_sha256':item['source_sha256'],'source_xlsx_row':rowno,'source_xlsx_sha256':sha(SOURCE)})

# Exact native-text table-row excerpt candidates for the ten GB/T 26610.5 Table 2/3 cells.
# Keep source typography/units/bounds; every change stays staging-only and requires visual signoff.
exact5=OUT/'research/2661045/table2_table3_exact_excerpt_candidates.csv'
if exact5.exists():
  exact_table_rows={
    144:'表2 面积后果等级划分\n单位为平方米\nA\nCA≤9.29',
    145:'表2 面积后果等级划分\n单位为平方米\nB\n9.29<CA≤92.9',
    146:'表2 面积后果等级划分\n单位为平方米\nC\n92.9<CA≤279',
    147:'表2 面积后果等级划分\n单位为平方米\nD\n279<CA≤929',
    148:'表2 面积后果等级划分\n单位为平方米\nE\nCA>929',
    149:'表3 热交换器管束的经济后果等级划分\n单位为元\nA\nCOF≤50 000',
    150:'表3 热交换器管束的经济后果等级划分\n单位为元\nB\n50 000<COF≤250 000',
    151:'表3 热交换器管束的经济后果等级划分\n单位为元\nC\n250 000<COF≤750 000',
    152:'表3 热交换器管束的经济后果等级划分\n单位为元\nD\n750 000<COF≤5 000 000',
    153:'表3 热交换器管束的经济后果等级划分\n单位为元\nE\nCOF>5 000 000',
  }
  with exact5.open(encoding='utf-8-sig',newline='') as _f:
    for item in csv.DictReader(_f):
      rowno=int(item['excel_row']); old=s(item['old_value']); new=exact_table_rows[rowno]
      assert s(ws.cell(rowno,6).value)==old, (rowno,ws.cell(rowno,6).value,old)
      assert len(new)<=300, (rowno,len(new))
      ws.cell(rowno,6).value=new
      detailed_bbox='grade='+item['grade_bbox']+'; numeric_spans='+item['source_numeric_spans']+'; page rendered at 2x in work/data_staging/p15/A_RBI/266105_pdfindex'+item['pdf_page_index_0based']+'.png'
      loc={'excel_row':rowno,'clause_id':item['clause_id'],'standard_code':'GB/T 26610.5-2022','clause_no':item['clause_id'].split('-2022-',1)[-1],'printed_page':item['printed_page'],'pdf_file':'upstream/01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf','pdf_page_index':item['pdf_page_index_0based'],'bbox':detailed_bbox,'table_id':item['table_id'],'table_row':item['table_row'],'table_column':item['table_column'],'element_name':'表题 + 单位 + 等级 + 数值范围','evidence_quote':new,'evidence_level':'PDF_RENDERED_VISUAL_CROSSCHECK__HUMAN_SIGNOFF_PENDING','review_required':'TRUE_HUMAN_VISUAL_CONFIRMATION','source_pdf_sha256':item['source_pdf_sha256']}
      source_ledger.append({**loc,'field':'原文摘录','quote':new,'quote_chars':len(new),'quote_mode':'TABLE_ROW_VISUALLY_TRANSCRIBED_CANDIDATE','quote_exactness':'VISUAL_CROSSCHECK_COMPLETE__HUMAN_SIGNOFF_PENDING','human_review':'NOT_SIGNED'})
      changelog.append({'excel_row':rowno,'clause_id':item['clause_id'],'field':'原文摘录','old_value':old,'new_value':new,'action':'REPLACE_TABLE_PARAPHRASE_WITH_SOURCE_ROW_EXCERPT','why':item['why'],'pdf_file':loc['pdf_file'],'printed_page':item['printed_page'],'pdf_page_index':item['pdf_page_index_0based'],'bbox':detailed_bbox,'table_id':item['table_id'],'table_row':item['table_row'],'table_column':item['table_column'],'element_name':loc['element_name'],'evidence_quote':new,'evidence_level':loc['evidence_level'],'review_required':'TRUE_HUMAN_VISUAL_CONFIRMATION','source_pdf_sha256':item['source_pdf_sha256'],'source_xlsx_row':rowno,'source_xlsx_sha256':sha(SOURCE)})

# Capture every current F row as an auditable source-quote queue with page/index evidence where available.
# Do not claim exactness merely from native search; textual/visual quotation alignment needs separate adjudication.
for excel_row,d in records:
    code=s(d.get('standard_code')); clause=s(d.get('clause_no'))
    paths={
      'GB/T 26610.1-2022':'GBT26610.1-2022_基本要求和实施程序.pdf',
      'GB/T 26610.2-2022':'GBT26610.2-2022_基于风险的检验策略.pdf',
      'GB/T 26610.3-2014':'GBT26610.3-2014_风险的定性分析方法.pdf',
      'GB/T 26610.4-2022':'GBT26610.4-2022_失效可能性定量分析方法.pdf',
      'GB/T 26610.5-2022':'GBT26610.5-2022_失效后果定量分析方法.pdf',
      'TSG 21-2016':'补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf'}
    rel='upstream/01_RBI标准/'+paths.get(code,'')
    evidence_pages=[]
    # Native-page cache page text can locate clause number but cannot prove table OCR/visual quote.
    needle=clause.replace('.','\\.')
    for obj in cache_rows:
      if obj.get('source_pdf','').endswith(paths.get(code,'__none__')) and re.search(r'(?<!\\d)'+needle+r'(?!\\d)',obj.get('text','')):
        evidence_pages.append(str(obj.get('pdf_page_index')))
        if len(evidence_pages)>=3: break
    row={'excel_row':excel_row,'clause_id':s(d['clause_id']),'standard_code':code,'clause_no':clause,'printed_page':d.get('page') or '', 'pdf_file':rel,'pdf_page_indexes_native_clause_match':'|'.join(evidence_pages),'field':'原文摘录','quote':s(d.get('原文摘录')),'quote_chars':len(s(d.get('原文摘录'))),'quote_mode':'TABLE_SUMMARY_PATTERN' if re.search(r'表\\s*\\d|表\\d|图1|对应.*等级|对应.*因子| → |→',s(d.get('原文摘录'))) else 'PROSE_OR_UNKNOWN','evidence_level':'E0_CLAUSE_PAGE_MACHINE_LOCATOR_ONLY' if evidence_pages else 'REVIEW_NEEDED_NO_NATIVE_PAGE_MATCH','quote_exactness':'NOT_ESTABLISHED','human_review':'NOT_SIGNED'}
    source_ledger.append(row)
    if not evidence_pages:
      unresolved.append({'excel_row':excel_row,'clause_id':s(d['clause_id']),'standard_code':code,'field':'F 原文摘录','issue':'native cache未定位到条款号；需视觉查原始PDF及具体表格单元，不判为错误','source_pdf':rel,'printed_page':d.get('page') or '', 'priority':'P0_TABLE_OR_NO_MATCH'})

# Replace incomplete cache-only locators with the detailed standard-specific table locator
# inventories (still E1/E2 pending; locator evidence is not an accuracy verdict).
locator_rows={}
loc4=OUT/'research/2661045/table_transcription_source_locator.csv'
if loc4.exists():
  with loc4.open(encoding='utf-8-sig',newline='') as _f:
    for item in csv.DictReader(_f):
      locator_rows[int(item['excel_row'])]=item
for item in locator_rows.values():
  rowno=int(item['excel_row'])
  unresolved=[u for u in unresolved if int(u['excel_row'])!=rowno]
  source_ledger.append({'excel_row':rowno,'clause_id':item['clause_id'],'standard_code':item['standard_code'],'clause_no':item['clause_no'],'printed_page':item['printed_page'],'pdf_file':item['source_pdf'],'pdf_page_index':item['pdf_page_index_0based'],'bbox':item['bbox'],'table_id':item['table_id'],'table_row':item['table_row_or_cell'],'table_column':'NOT_MACHINE_RESOLVED','field':'原文摘录','quote':item['current_F_quote'],'quote_chars':len(item['current_F_quote']),'quote_mode':'TABLE_SUMMARY_CANDIDATE','evidence_quote':item['native_page_evidence_excerpt'],'evidence_level':item['evidence_level'],'quote_exactness':'NOT_ESTABLISHED','human_review':item['manual_review_status'],'review_required':item['review_required'],'source_pdf_sha256':item['source_pdf_sha256']})

loc3=OUT/'research/266103/266103_row_evidence_index.csv'
if loc3.exists():
  with loc3.open(encoding='utf-8-sig',newline='') as _f:
    for item in csv.DictReader(_f):
      rowno=int(item['excel_row'])
      unresolved=[u for u in unresolved if int(u['excel_row'])!=rowno]
      source_ledger.append({'excel_row':rowno,'clause_id':item['clause_id'],'standard_code':'GB/T 26610.3-2014','clause_no':item['clause_no'],'printed_page':item['printed_page_source_evidence'],'pdf_file':item['pdf_file'],'pdf_page_index':item['pdf_page_index_0based'],'bbox':item['bbox'],'table_id':item['table_id'],'table_row':item['table_row_selector'],'table_column':item['table_column_selector'],'field':'原文摘录 / G-J / T','quote':item['current_content_raw'],'quote_chars':len(item['current_content_raw']),'quote_mode':'FIGURE_OR_TABLE_LOCATOR_ROW','evidence_quote':item['native_page_text_excerpt'],'evidence_level':item['evidence_level'],'quote_exactness':'ONLY_FIG1_25_CELLS_VISUALLY_MATCHED; OTHER_TABLES_PENDING','human_review':item['reviewer'],'review_required':item['human_review_required'],'source_pdf_sha256':sha(PDF)})

# Keep the conditional table work rows located even when the PDF text encoding blocks extraction.
loc2=OUT/'research/266102_tsg/review_scope_rows.jsonl'
if loc2.exists():
  import json as _json
  with loc2.open(encoding='utf-8') as _f:
    for line in _f:
      try: item=_json.loads(line)
      except Exception: continue
      source_ledger.append({'excel_row':item['excel_row'],'clause_id':item['clause_id'],'standard_code':item['standard_code'],'clause_no':item['clause_no'],'printed_page':item['printed_page'],'pdf_file':'upstream/01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf' if item['standard_code']=='GB/T 26610.2-2022' else 'upstream/01_RBI标准/补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf','pdf_page_index':'','bbox':'not resolved','table_id':'','table_row':'','table_column':'','field':'F/K/L/T','quote':item['current_F'],'quote_chars':len(item['current_F']),'quote_mode':'CONDITIONAL_INSPECTION_REVIEW_SCOPE','evidence_quote':'条款范围样本；本机子审计摘要详见 research/266102_tsg/README.md','evidence_level':'CLAUSE_LOCATOR_PENDING_PAGE_EXACTNESS','human_review':'NOT_SIGNED','review_required':'TRUE'})

# Complete disposition for all 221 heuristic table-transcription candidates.
table_queue=ROOT/'work/data_staging/p14/agents/B/rbi_table_quote_review_queue.jsonl'
table_coverage=[]
if table_queue.exists():
  import json as _json
  loc3_by_row={int(x['excel_row']):x for x in csv.DictReader((OUT/'research/266103/266103_row_evidence_index.csv').open(encoding='utf-8-sig',newline=''))} if loc3.exists() else {}
  exact5_by_row={int(x['excel_row']):x for x in csv.DictReader((OUT/'research/2661045/table2_table3_exact_excerpt_candidates.csv').open(encoding='utf-8-sig',newline=''))} if exact5.exists() else {}
  for line in table_queue.open(encoding='utf-8'):
    item=_json.loads(line); rowno=int(item['excel_row']); std=item['standard_code']; pnum=int(item['printed_page']) if item.get('printed_page') else None
    pdf_name={'GB/T 26610.2-2022':'upstream/01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf','GB/T 26610.3-2014':'upstream/01_RBI标准/GBT26610.3-2014_风险的定性分析方法.pdf','GB/T 26610.4-2022':'upstream/01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf','GB/T 26610.5-2022':'upstream/01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf'}[std]
    index=''; bbox=''; status='LOCATED_PAGE_TABLE_CELL_PENDING'; method='P14 heuristic source + printed page; table geometry not manually checked'; detail='table/cell contents and F exactness remain REVIEW_NEEDED'
    if std=='GB/T 26610.3-2014' and rowno in loc3_by_row:
      x=loc3_by_row[rowno]; index=x['pdf_page_index_0based']; bbox=x['bbox']; status='LOCATED_PAGE_CELL_PENDING'; method=x['evidence_level']; detail=x['source_locator']+'; '+x['T_class']
      if 'VISUAL_CELL_MATCH' in x['matrix_check']: status='FIG1_VISUAL_MATRIX_MATCHED_HUMAN_SIGNOFF_PENDING'
    elif std in ('GB/T 26610.4-2022','GB/T 26610.5-2022') and rowno in locator_rows:
      x=locator_rows[rowno]; index=x['pdf_page_index_0based']; bbox=x['bbox']; method=x['evidence_level']; status='LOCATED_PAGE_CELL_PENDING' if x['bbox']=='NOT_MACHINE_RESOLVED' else 'LOCATED_CELL_SPAN_VISUAL_CONFIRMATION_PENDING'; detail=x['native_page_evidence_excerpt']
      if std=='GB/T 26610.5-2022' and rowno in exact5_by_row:
        status='VISUAL_SOURCE_CROSSCHECK_HUMAN_SIGNOFF_PENDING'; detail=exact5_by_row[rowno]['proposed_exact_excerpt']
    elif std=='GB/T 26610.2-2022' and pnum is not None:
      index=str(pnum+4); bbox='NOT_MACHINE_RESOLVED'; status='PHYSICAL_PAGE_SEQUENCE_MAPPED_CELL_UNRESOLVED'; method='PDF pagination offset +4 (sample validated at printed p3/index7); target page not individually rendered'; detail='P14 heuristic only; figure/table structure and cell reading order need visual check'
    table_coverage.append({'excel_row':rowno,'clause_id':item['candidate_id'],'standard_code':std,'clause_no':item['clause_no'],'printed_page':pnum or '', 'pdf_file':pdf_name,'pdf_page_index':index,'table_id':item['clause_no'] if item['clause_no'].startswith('表') else '', 'source_locator':detail,'bbox_or_failure':bbox,'current_F_quote':item['raw_quote'],'review_status':status,'method':method,'evidence_level':'MACHINE_LOCATOR_NOT_ACCURACY_VERDICT','reviewer_signature':'EMPTY'})

wb.save(CANDIDATE)

# Flat CSV exports.
def write_csv(path,rows,fields):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore'); w.writeheader(); w.writerows(rows)
write_csv(TRIAGE,triage,['excel_row','clause_id','standard_code','clause_no','field','value_state','field_class','reason','draft_status','source_pdf','printed_page','pdf_page_index','table_id','table_row','table_column','evidence_quote','machine_review'])
write_csv(CHANGELOG,changelog,['excel_row','clause_id','field','old_value','new_value','action','why','pdf_file','printed_page','pdf_page_index','bbox','table_id','table_row','table_column','element_name','evidence_quote','evidence_level','review_required','source_pdf_sha256','source_xlsx_row','source_xlsx_sha256'])
write_csv(SOURCE_LEDGER,source_ledger,['excel_row','clause_id','standard_code','clause_no','printed_page','pdf_file','pdf_page_indexes_native_clause_match','field','quote','quote_chars','quote_mode','evidence_level','quote_exactness','human_review','pdf_page_index','bbox','table_id','table_row','table_column','element_name','evidence_quote','review_required','source_pdf_sha256'])
write_csv(UNRESOLVED,unresolved,['excel_row','clause_id','standard_code','field','issue','source_pdf','printed_page','priority'])
write_csv(OUT/'P15_A_TABLE_REVIEW_COVERAGE.csv',table_coverage,['excel_row','clause_id','standard_code','clause_no','printed_page','pdf_file','pdf_page_index','table_id','source_locator','bbox_or_failure','current_F_quote','review_status','method','evidence_level','reviewer_signature'])
count_rows=[]
for code in sorted({s(d.get('standard_code')) for _,d in records}):
  subset=[(r,d) for r,d in records if s(d.get('standard_code'))==code]
  for field in headers:
    src=sum(1 for _,d in subset if d.get(field) is not None and s(d.get(field)))
    cand=sum(1 for excel_row,_ in subset if ws.cell(excel_row,headers.index(field)+1).value is not None and s(ws.cell(excel_row,headers.index(field)+1).value))
    count_rows.append({'standard_code':code,'field':field,'records':len(subset),'source_nonempty':src,'candidate_nonempty':cand,'delta':cand-src,'supported_recall':'N/A_SOURCE_SUPPORT_NOT_FULLY_ADJUDICATED'})
write_csv(OUT/'P15_A_FIELD_COUNTS.csv',count_rows,['standard_code','field','records','source_nonempty','candidate_nonempty','delta','supported_recall'])

summary={'run_utc':datetime.now(timezone.utc).isoformat(),'source_xlsx':str(SOURCE.relative_to(ROOT)),'source_xlsx_sha256':sha(SOURCE),'source_pdf':str(PDF.relative_to(ROOT)),'source_pdf_sha256':sha(PDF),'source_rows':len(records),'source_columns':len(headers),'candidate':str(CANDIDATE.relative_to(ROOT)),'candidate_sha256':sha(CANDIDATE),'candidate_data_changes':len(changelog),'cells_changed':sum(1 for x in changelog if x['old_value']!=x['new_value']),'matrix_rows_checked':len(mat),'matrix_values_checked':25,'field_triage_cells':len(triage),'source_quote_rows':len(records),'table_heuristic_rows_total':len(table_coverage),'table_locator_rows_by_status':dict(__import__('collections').Counter(x['review_status'] for x in table_coverage)),'native_clause_locator_unresolved':len(unresolved),'machine_quote_exactness_claimed':0,'reviewer_signatures_added':0,'status':'REVIEW_BLOCKED'}
(OUT/'P15_A_RUN_SUMMARY.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))

