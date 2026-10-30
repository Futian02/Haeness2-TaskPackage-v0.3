"""P15 deterministic human-audit pilot and guarded staging-only merge helper."""
from __future__ import annotations
import argparse, csv, hashlib, json, random, shutil, subprocess, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
UP=ROOT/"upstream"
MERGED=ROOT/"work/data_staging/p15/merged"
OFFICIAL=UP/"05_提交区/clauses.jsonl"
CONVERT=UP/"04_接口文档/convert.py"
VALIDATE=UP/"04_接口文档/validate.py"
RBI_BASE=UP/"01_RBI标准/条款摘录表_RBI.xlsx"
DM_BASE=UP/"02_损伤机理参考/条款摘录表_损伤机理.xlsx"
SHEETS=["填写说明","摘录表","填写示例","枚举值"]
PDF_MAP={
 "GB/T 26610.1-2022":"01_RBI标准/GBT26610.1-2022_基本要求和实施程序.pdf",
 "GB/T 26610.2-2022":"01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf",
 "GB/T 26610.3-2014":"01_RBI标准/GBT26610.3-2014_风险的定性分析方法.pdf",
 "GB/T 26610.4-2022":"01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf",
 "GB/T 26610.5-2022":"01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf",
 "TSG 21-2016":"01_RBI标准/补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf",
 "GB/T 30579-2022":"02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf",
}

def read_jsonl(path):
 return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]

def sha(path:Path)->str:
 h=hashlib.sha256()
 with path.open("rb") as f:
  for block in iter(lambda:f.read(1024*1024),b""): h.update(block)
 return h.hexdigest()

def seed50(output:Path, seed:int=20261029, source:Path=OFFICIAL):
 rows=read_jsonl(source)
 rng=random.Random(seed); chosen=sorted(rng.sample(rows,min(50,len(rows))),key=lambda x:x["clause_id"])
 fields=["sample_no","seed","clause_id","standard_code","clause_no","title","source_pdf","printed_page","pdf_page_index","bbox_or_table_cell","evidence_quote","conditions_to_verify","extractor_identity_check","independent_reviewer_identity","human_status","pass_fail","error_type","review_reason","review_timestamp"]
 output.parent.mkdir(parents=True,exist_ok=True)
 with output.open("w",encoding="utf-8-sig",newline="") as f:
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
  for i,r in enumerate(chosen,1):
   standard=r.get("standard_code",""); pdf_rel=("upstream/"+PDF_MAP[standard]) if standard in PDF_MAP else ""
   cond=r.get("applicable_conditions") or {}
   cond_text=json.dumps(cond,ensure_ascii=False,sort_keys=True)
   w.writerow({"sample_no":i,"seed":seed,"clause_id":r.get("clause_id"),"standard_code":standard,"clause_no":r.get("clause_no"),"title":r.get("title"),"source_pdf":pdf_rel,"printed_page":r.get("page",""),"pdf_page_index":"","bbox_or_table_cell":"","evidence_quote":r.get("content_raw",""),"conditions_to_verify":cond_text,"extractor_identity_check":"PENDING_REAL_IDENTITY","independent_reviewer_identity":"","human_status":"PENDING_HUMAN_REVIEW","pass_fail":"","error_type":"","review_reason":"","review_timestamp":""})
 print(f"human_review_rows={len(chosen)} seed={seed} pending={len(chosen)} signed=0 source={source} path={output}")
 return 0

def within(path:Path, base:Path)->bool:
 try: path.resolve().relative_to(base.resolve()); return True
 except ValueError: return False

def norm(v): return "" if v is None else str(v).strip()

def workbook_records(path:Path):
 wb=load_workbook(path,read_only=True,data_only=True)
 try:
  if wb.sheetnames!=SHEETS: raise ValueError(f"workbook sheets changed: {path} {wb.sheetnames}")
  ws=wb["摘录表"]; it=ws.iter_rows(values_only=True); headers=[norm(v) for v in next(it)]
  rec=[]
  for excel_row,values in enumerate(it,2):
   row=dict(zip(headers,values))
   if any(norm(v) for v in row.values()): rec.append((excel_row,row))
  return headers,rec
 finally: wb.close()

def compare_candidate(base:Path,candidate:Path,ledger:Path,kind:str):
 base_headers,old_rows=workbook_records(base); cand_headers,new_rows=workbook_records(candidate)
 if base_headers!=cand_headers: raise ValueError(f"header mismatch {candidate}")
 if [norm(r.get("clause_id")) for _,r in old_rows]!=[norm(r.get("clause_id")) for _,r in new_rows]: raise ValueError(f"row order/primary keys changed in {candidate}")
 with ledger.open(encoding="utf-8-sig",newline="") as f: logs=list(csv.DictReader(f))
 log_by_key={}
 for row in logs:
  key=(int(row.get("excel_row") or 0),norm(row.get("field")))
  if key in log_by_key: raise ValueError(f"duplicate changelog key {ledger} {key}")
  log_by_key[key]=row
 diffs=[]; seen=set()
 for (rownum,old),(candidate_rownum,new) in zip(old_rows,new_rows):
  assert rownum==candidate_rownum
  cid=norm(new.get("clause_id"))
  for field in base_headers:
   if norm(old.get(field))==norm(new.get(field)): continue
   key=(rownum,field); seen.add(key); log=log_by_key.get(key)
   if not log: raise ValueError(f"unlogged cell change {kind} row={rownum} clause={cid} field={field}")
   if norm(log.get("old_value"))!=norm(old.get(field)) or norm(log.get("new_value"))!=norm(new.get(field)):
    raise ValueError(f"changelog old/new mismatch {kind} row={rownum} clause={cid} field={field}")
   diffs.append({"kind":kind,"row":rownum,"clause_id":cid,"field":field,"old_value":norm(old.get(field)),"new_value":norm(new.get(field)),"action":norm(log.get("action")),"source_pdf":norm(log.get("pdf_file") or log.get("standard_code")),"printed_page":norm(log.get("printed_page")),"pdf_page_index":norm(log.get("pdf_page_index")),"bbox_or_locator":norm(log.get("bbox") or log.get("evidence_locator")),"table_id":norm(log.get("table_id")),"table_row":norm(log.get("table_row")),"table_column":norm(log.get("table_column")),"evidence_quote":norm(log.get("evidence_quote")),"evidence_level":norm(log.get("evidence_level") or log.get("method")),"reviewer_status":norm(log.get("review_required") or log.get("reviewer_status")),"human_reviewer":norm(log.get("human_reviewer"))})
 extra=set(log_by_key)-seen
 if extra: raise ValueError(f"changelog records without xlsx diff {kind}: {sorted(extra)[:5]} total={len(extra)}")
 print(f"{kind}: rows={len(old_rows)} field_changes={len(diffs)} changelog_rows={len(logs)} exact_match=True")
 return diffs

def make_convert_eligible(src:Path,dest:Path):
 shutil.copy2(src,dest)
 wb=load_workbook(dest)
 ws=wb["摘录表"]; headers=[norm(x.value) for x in ws[1]]
 if "抽取人" not in headers: raise ValueError("RBI candidate has no extractor field")
 col=headers.index("抽取人")+1; idcol=headers.index("clause_id")+1
 remove=[]
 for row in range(2,ws.max_row+1):
  if norm(ws.cell(row,col).value)=="" and norm(ws.cell(row,idcol).value): remove.append(row)
 for row in reversed(remove): ws.delete_rows(row,1)
 wb.save(dest);wb.close()
 _,eligible=workbook_records(dest)
 print(f"convert_eligible_RBI_rows={len(eligible)} excluded_unsigned_rows={len(remove)}; source preview remains intact")
 return len(remove)

def write_candidate_tool_probe(candidate_clauses:Path,log_path:Path):
 root=MERGED/"candidate_tool_probe"; submit=root/"05_提交区"; api=root/"04_接口文档"
 submit.mkdir(parents=True,exist_ok=True); (api/"schema").mkdir(parents=True,exist_ok=True)
 originals=[(UP/"05_提交区/tools.py",submit/"tools.py"),(UP/"05_提交区/_selftest_tools.py",submit/"_selftest_tools.py"),(UP/"05_提交区/_df_rules.py",submit/"_df_rules.py"),(UP/"04_接口文档/enums.json",api/"enums.json"),(UP/"04_接口文档/material_family.json",api/"material_family.json"),(UP/"04_接口文档/risk_level_mapping.xlsx",api/"risk_level_mapping.xlsx")]
 for src,dst in originals:
  if src.exists(): shutil.copy2(src,dst)
 shutil.copy2(candidate_clauses,submit/"clauses.jsonl")
 shutil.copy2(UP/"04_接口文档/schema/tool_response.schema.json",api/"schema/tool_response.schema.json")
 if sha(submit/"tools.py")!=sha(UP/"05_提交区/tools.py"): raise ValueError("candidate tool probe code differs from frozen teacher tools.py")
 cmd=[sys.executable,"-B",str(submit/"_selftest_tools.py")]
 proc=subprocess.run(cmd,cwd=root,text=True,capture_output=True,encoding="utf-8",errors="replace")
 log_path.write_text("CMD="+subprocess.list2cmdline(cmd)+"\nEXIT="+str(proc.returncode)+"\n"+proc.stdout+proc.stderr,encoding="utf-8")
 print(proc.stdout,end="");print(proc.stderr,end="",file=sys.stderr)
 print(f"candidate_three_tool_selftest_exit={proc.returncode}; frozen_tools_sha256={sha(submit/'tools.py')}; log={log_path}")
 return proc.returncode,submit

def candidate_semantic_audit(path:Path):
 rows=read_jsonl(path); by={r["clause_id"]:r for r in rows}; checks=[]
 expected_matrix={(str(n),letter) for n in range(1,6) for letter in "ABCDE"}
 matrix=[r for r in rows if r.get("clause_id","").startswith("GBT26610.3-2014-7-")]
 observed={(str((r.get("rbi_detail") or {}).get("likelihood")),str((r.get("rbi_detail") or {}).get("consequence"))) for r in matrix}
 checks.append({"check":"matrix_5x5_completeness","pass":len(matrix)==25 and observed==expected_matrix,"observed_cells":len(matrix),"expected_cells":25,"scope":"syntactic completeness; mapping crosschecked separately by field gate"})
 area={"a":"CA≤9.29","b":"9.29<CA≤92.9","c":"92.9<CA≤279","d":"279<CA≤929","e":"CA>929"}
 for suffix,needle in area.items():
  cid=f"GBT26610.5-2022-表2-{suffix}"; record=by.get(cid); content=(record or {}).get("content_raw","")
  checks.append({"check":"area_consequence_boundary_"+suffix,"pass":bool(record and needle in content),"clause_id":cid,"expected_fragment":needle,"observed_content":content})
 dm={"GBT30579-2022-5.23":{"temp":[None,None],"need_other":"49℃～66℃"},"GBT30579-2022-5.6":{"temp":[177,427],"need_other":"218℃～400℃"},"GBT30579-2022-6.7":{"pressure":[None,None],"need_other":"0.0003 MPa"}}
 for cid,expected in dm.items():
  record=by.get(cid) or {}; cond=record.get("applicable_conditions",{}); ok=True
  if "temp" in expected: ok &= cond.get("temperature_c")==expected["temp"]
  if "pressure" in expected: ok &= cond.get("pressure_mpa")==expected["pressure"]
  ok &= expected["need_other"] in str(cond.get("other",""))
  checks.append({"check":"damage_semantics_"+cid,"pass":bool(ok),"observed_conditions":cond,"scope":"machine consistency only; independent visual/human verification pending"})
 result={"candidate_clause_count":len(rows),"checks":checks,"machine_checks_pass":all(x["pass"] for x in checks),"human_pdf_fidelity":"NOT_VERIFIED"}
 out=MERGED/"P15_C_CANDIDATE_SEMANTIC_AUDIT.json";out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print(f"candidate_semantics checks={len(checks)} pass={sum(c['pass'] for c in checks)}/{len(checks)} human_pdf_fidelity=NOT_VERIFIED")
 return result

def write_manifest(candidate_jsonl:Path):
 files={"source_rbi_xlsx":RBI_BASE,"source_dm_xlsx":DM_BASE,"official_clauses":OFFICIAL,"candidate_rbi_preview":MERGED/"条款摘录表_RBI_P15_A_preview_full_unreleased.xlsx","candidate_rbi_convert_eligible":MERGED/"条款摘录表_RBI_P15_A_convert_eligible.xlsx","candidate_dm":MERGED/"条款摘录表_损伤机理_P15_B_candidate.xlsx","candidate_clauses":candidate_jsonl,"teacher_converter":CONVERT,"teacher_validator":VALIDATE,"teacher_tools":UP/"05_提交区/tools.py"}
 manifest={k:{"path":str(v),"sha256":sha(v)} for k,v in files.items()}
 (MERGED/"P15_C_SHA256_MANIFEST.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 return manifest

def merge(rbi:Path,dm:Path):
 MERGED.mkdir(parents=True,exist_ok=True)
 targets=[("条款摘录表_RBI_P15_A_preview_full_unreleased.xlsx",rbi),("条款摘录表_损伤机理_P15_B_candidate.xlsx",dm)]
 for target,src in targets:
  src=src.resolve()
  if not src.is_file(): raise FileNotFoundError(src)
  # Inputs are copied byte-for-byte; source candidate paths are never opened for writing.
  shutil.copy2(src,MERGED/target)
 a_changes=compare_candidate(RBI_BASE,rbi,ROOT/"work/data_staging/p15/A_RBI/P15_A_FIELD_CHANGELOG.csv","RBI")
 b_changes=compare_candidate(DM_BASE,dm,ROOT/"docs/data_quality/P15_B_FIELD_CHANGELOG.csv","DM")
 with (MERGED/"DIFF_LEDGER.csv").open("w",encoding="utf-8-sig",newline="") as f:
  fields=list((a_changes+b_changes)[0]) if a_changes or b_changes else ["kind","row","clause_id","field","old_value","new_value","action","source_pdf","printed_page","pdf_page_index","bbox_or_locator","table_id","table_row","table_column","evidence_quote","evidence_level","reviewer_status","human_reviewer"]
  w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(a_changes+b_changes)
 unsigned_count=make_convert_eligible(MERGED/targets[0][0],MERGED/"条款摘录表_RBI_P15_A_convert_eligible.xlsx")
 eligible_rbi=MERGED/"条款摘录表_RBI_P15_A_convert_eligible.xlsx"
 eligible_dm=MERGED/targets[1][0]
 out=MERGED/"clauses_candidate.jsonl"
 cmd=[sys.executable,"-B",str(CONVERT),"--inputs",str(eligible_rbi),str(eligible_dm),"--output",str(out)]
 print("RUN",subprocess.list2cmdline(cmd)); res=subprocess.run(cmd,cwd=UP,text=True,capture_output=True,encoding="utf-8",errors="replace");print(res.stdout,end="");print(res.stderr,end="",file=sys.stderr)
 if res.returncode: return res.returncode
 validate=[sys.executable,"-B",str(VALIDATE),"--final","--clauses",str(out),"--trajectories",str(UP/"05_提交区/trajectories.jsonl")]
 print("RUN",subprocess.list2cmdline(validate)); res=subprocess.run(validate,cwd=UP,text=True,capture_output=True,encoding="utf-8",errors="replace");print(res.stdout,end="");print(res.stderr,end="",file=sys.stderr)
 if res.returncode: return res.returncode
 official={r["clause_id"]:r for r in read_jsonl(OFFICIAL)}; candidate={r["clause_id"]:r for r in read_jsonl(out)}
 ids=sorted(set(official)|set(candidate)); changes=[]
 for cid in ids:
  old,new=official.get(cid),candidate.get(cid)
  if old is None: action="added"
  elif new is None: action="removed"
  elif old==new: action="unchanged"
  else: action="modified"
  diff=[]
  if old and new:
   for k in sorted(set(old)|set(new)):
    if old.get(k)!=new.get(k): diff.append(k)
  if old is None or new is None:
   changes.append({"clause_id":cid,"action":action,"field":"__record__","old_value":"","new_value":"","source_evidence":"no row-level evidence linked"})
  elif old!=new:
   for field in sorted(set(old)|set(new)):
    if old.get(field)!=new.get(field): changes.append({"clause_id":cid,"action":"modified","field":field,"old_value":json.dumps(old.get(field),ensure_ascii=False),"new_value":json.dumps(new.get(field),ensure_ascii=False),"source_evidence":"see corresponding A/B field ledger; human review still pending"})
 record_counts={a:sum(1 for cid in set(official)|set(candidate) if ("added" if cid not in official else "removed" if cid not in candidate else "modified" if official[cid]!=candidate[cid] else "unchanged")==a) for a in ("added","modified","unchanged","removed")}
 with (MERGED/"P15_C_CANDIDATE_JSONL_DIFF.csv").open("w",encoding="utf-8-sig",newline="") as f:
  fields=["clause_id","action","field","old_value","new_value","source_evidence"];w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(changes)
 print("CANDIDATE_JSONL_RECORD_COUNTS",json.dumps(record_counts,ensure_ascii=False),"changed_field_rows",len(changes),"unsigned_draft_rows_excluded",unsigned_count)
 print("NOTE source evidence, author authenticity, and PDF fidelity remain an independent gate; converter PASS does not establish them.")
 candidate_semantic_audit(out)
 write_manifest(out)
 tool_code,submit=write_candidate_tool_probe(out,MERGED/"P15_C_CANDIDATE_TOOL_SELFTEST.log")
 if tool_code: return tool_code
 return 0

if __name__=="__main__":
 p=argparse.ArgumentParser(); sub=p.add_subparsers(dest="cmd",required=True)
 s=sub.add_parser("sample50");s.add_argument("--out",type=Path,default=ROOT/"work/data_staging/p15/C_QA/P15_C_HUMAN_REVIEW50.csv");s.add_argument("--seed",type=int,default=20261029);s.add_argument("--source",type=Path,default=OFFICIAL)
 m=sub.add_parser("merge");m.add_argument("--rbi",type=Path,required=True);m.add_argument("--dm",type=Path,required=True)
 a=p.parse_args()
 try: raise SystemExit(seed50(a.out,a.seed,a.source) if a.cmd=="sample50" else merge(a.rbi,a.dm))
 except Exception as e: print(f"P15 staging operation failed: {type(e).__name__}: {e}",file=sys.stderr);raise SystemExit(2)
