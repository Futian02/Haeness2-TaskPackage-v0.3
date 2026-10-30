"""Verify the R8-A in-place draft update without touching teacher data sources."""
from __future__ import annotations
import csv, hashlib, random, zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
X=ROOT/'upstream'/'01_RBI标准'/'条款摘录表_RBI.xlsx'
B=X.with_name(X.name+'.bak_owner_before_rbi1_20261028')
A=ROOT/'docs'/'audits'
FIELDS=list('ABCDEFGHIJKLMNOPQRSTUVW')
PART='xl/worksheets/sheet2.xml'

def sha(b): return hashlib.sha256(b).hexdigest()
def check(ok,msg):
    if not ok: raise AssertionError(msg)

def main():
    before=load_workbook(BytesIO(B.read_bytes()),data_only=False,read_only=False)
    after=load_workbook(X,data_only=False,read_only=False)
    check(before.sheetnames==after.sheetnames,'sheetname/order changed')
    check(before.sheetnames==['填写说明','摘录表','填写示例','枚举值'],f'unexpected sheets {before.sheetnames}')
    # Compare all unaffected sheets at cell/value/style/formula and layout levels.
    unaffected=[]
    for name in before.sheetnames:
        if name=='摘录表': continue
        a,b=before[name],after[name]
        check(a.max_row==b.max_row and a.max_column==b.max_column,f'{name} dimension changed')
        for row in range(1,a.max_row+1):
            for col in range(1,a.max_column+1):
                ca,cb=a.cell(row,col),b.cell(row,col)
                check((ca.value,ca.data_type,ca.style_id)==(cb.value,cb.data_type,cb.style_id),f'{name}!{ca.coordinate} changed')
        check(list(a.merged_cells.ranges)==list(b.merged_cells.ranges),f'{name} merges changed')
        check(str(a.data_validations)==str(b.data_validations),f'{name} validation changed')
        check(a.freeze_panes==b.freeze_panes,f'{name} freeze panes changed')
        unaffected.append(name)
    # Existing data/style must remain byte-for-byte semantically identical through row 473.
    a,b=before['摘录表'],after['摘录表']
    for row in range(1,474):
        for col in range(1,24):
            x,y=a.cell(row,col),b.cell(row,col)
            check((x.value,x.data_type,x.style_id)==(y.value,y.data_type,y.style_id),f'old cell {x.coordinate} changed')
    rows=[]
    for r in range(474,487):
        rec={c:b[f'{c}{r}'].value for c in FIELDS}
        check(rec['A'] and rec['F'] and rec['W'] and 'DRAFT' in str(rec['W']),f'row {r} not draft-complete')
        check(rec['U'] in (None,'') and rec['V'] in (None,''),f'row {r} fabricated signer')
        check(all(rec[c] in (None,'') for c in 'GHIJKLMNOPQRST'),f'row {r} has unsupported classification/condition')
        check(len(str(rec['F']))<=300,f'row {r} quote >300')
        rows.append(rec)
    ids=[r['A'] for r in rows]
    check(len(ids)==13 and len(set(ids))==13,'13 candidate IDs missing/duplicated')
    check(len({r[0].value for r in b.iter_rows(min_row=2,max_row=473,min_col=1,max_col=1)})==472,'old IDs not unique')
    # The only changed XLSX package member must be the target data worksheet XML.
    with zipfile.ZipFile(B) as za, zipfile.ZipFile(X) as zb:
        check(za.namelist()==zb.namelist(),'OOXML package part list changed')
        changed=[n for n in za.namelist() if za.read(n)!=zb.read(n)]
        check(changed==[PART],f'unexpected OOXML parts changed: {changed}')
    # Per-cell audit deltas, not an inferred diff summary.
    pdf_sha='af4a2b2aebe3b792635e9e14aff6b4255bcc0e20627c59bb656fae1721522c57'
    delta=[]
    for r in range(474,487):
        for col in FIELDS:
            old,new=a[f'{col}{r}'].value,b[f'{col}{r}'].value
            if old!=new:
                delta.append({'cell':f'{col}{r}','old':old or '','new':new or '',
                    'pdf_sha':pdf_sha,'printed_page':b[f'D{r}'].value or '',
                    'pdf_page':(int(b[f'D{r}'].value)+4) if b[f'D{r}'].value else '',
                    'confidence_source':'P2 candidate machine check plus previously recorded visual page review; no independent signer',
                    'source_excerpt':'exact quote is in F cell of this same row; see candidate clause mapping ledger',
                    'needs_human_review':'true'})
    with (A/'R8A_RBI_CELL_DELTA.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(delta[0]));w.writeheader();w.writerows(delta)
    # Deterministic 20-item numeric/table/formula sample; this is a human visual queue, not a passed review.
    pages=list(csv.DictReader((A/'R8A_RBI_PAGE_INVENTORY.csv').open(encoding='utf-8-sig',newline='')))
    candidates=[p for p in pages if p['section_table_formula_markers'] and any(k in p['section_table_formula_markers'] for k in ('表','式','公式','图'))]
    random.Random(8).shuffle(candidates)
    selected=candidates[:20]
    check(len(selected)==20,f'only {len(selected)} table/formula candidate pages found')
    with (A/'R8A_RBI_NUMERIC_FORMULA_RECHECK.csv').open('w',encoding='utf-8-sig',newline='') as f:
        fields=['standard_code','source_pdf','source_sha256','pdf_page_0based','printed_page','detected_items','review_status','reason']
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
        for p in selected:
            item={k:p[k] for k in fields[:5]}
            item['detected_items']=p['section_table_formula_markers']
            item.update(review_status='needs_human_visual_review',reason='OCR/page-marker discovery only; not numerically certified')
            w.writerow(item)
    # Summary stats and hashes are also written for the owner.
    with (A/'R8A_RBI_PAGE_INVENTORY.csv').open(encoding='utf-8-sig',newline='') as f: page_rows=list(csv.DictReader(f))
    with (A/'R8A_RBI_PDF_COVERAGE_LEDGER.csv').open(encoding='utf-8-sig',newline='') as f: clause_rows=list(csv.DictReader(f))
    text=['# R8-A RBI PDF → 原表审计','', '**状态：READY_FOR_OWNER_REVIEW；全页OCR索引已完成，但461页人工视觉/数值复核和真实双人签署仍未完成。新增13行仅DRAFT。**','',
      '## 原表与回滚','',f'- 原件备份：`{B}`；原SHA-256：`{sha(B.read_bytes())}`。',f'- 更新后原路径：`{X}`；SHA-256：`{sha(X.read_bytes())}`。',
      f'- 新增行：474–486；RBI源表条款从472变485；U/V签名均空；13条G–S/K等无来源字段留空。',
      f'- 23列、4工作表；`摘录表`既有A1:W473逐格值/类型/style_id一致；另3表值/公式/style/尺寸/合并/验证一致。',
      f'- OOXML ZIP仅`{PART}`变化；其余所有包成员字节相同；无格式重绘。',
      '', '## 六份PDF页面覆盖索引','',f'- 页级索引：{len(page_rows)}页；对应章节/既有条款映射CSV：{len(clause_rows)}行。OCR覆盖数：{sum(1 for x in page_rows if x["ocr_characters"])}；页面状态：{dict(Counter(x["review_status"] for x in page_rows))}。',
      '- OCR用于定位章节/表/公式，不能证明原文、印刷页、数值、单位、适用条件准确。逐页图像级复核仍需人工；目录/表格/公式和乱码页全部留待审核，不标“核验通过”。',
      '- 20项固定随机种子数值/表格/公式抽查清单位于`R8A_RBI_NUMERIC_FORMULA_RECHECK.csv`，全部`needs_human_visual_review`，没有伪称数值通过。',
      '', '## 当前转换和官方数据状态','',
      '- 真实抽取人与独立复核人缺失；老师`convert.py`按必填U字段门禁预期BLOCKED，不能合法将候选纳入JSONL。未改老师程序，未改正式`clauses.jsonl`。',
      '- `validate.py --final`核验的是未变更的既有正式JSONL，不是新草稿。',
      '- 13条映射及逐格变化见`R8A_RBI_PDF_COVERAGE_LEDGER.csv`、`R8A_RBI_CELL_DELTA.csv`；PDF哈希/页面数量见`R8A_RBI_PAGE_INVENTORY.csv`。',
      '', '## 数值、风险与周期门禁','',
      '- 26610.1的13条均只摘方法性原文，G/H/I/J/K及无依据条件字段留空；不从标准方法描述推导设备风险。',
      '- 已有25格风险矩阵数据未修改；未从OCR/示例量扩展风险或周期。TSG21安全状况1–5级不映射到L1–L4。',
      '- 表格/公式/周期数字抽查当前未获独立视觉复核，明确为待审核；不将OCR数字写入Excel。',
      '', '## 失败/阻塞','',
      '- 初次`r8a_build_coverage.py`已完整OCR 461页并写出两份索引，但在打开非`.xlsx`后缀备份时openpyxl拒绝扩展名；脚本已修复为BytesIO读取。索引数据已实际生成；本轮没有为脚本改动重复跑九分钟OCR。',
      '- 没有真人签名，因此正式转换BLOCKED；保留新增草稿，不伪造签核。无真人复核的461页也不能宣称已穷尽核证。',
      '- 不执行`validate.py`通过来冒充草稿验收；只对原始官方JSONL作回归。',
      '', '## 回滚','',
      f'- 回滚前先Owner核对两份SHA；可由Owner批准后将备份恢复原路径：`{B}`。本轮未提交/部署。']
    (A/'R8A_RBI_FINAL_AUDIT.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    before.close(); after.close()
    print('WORKBOOK_STRUCTURE=PASS')
    print('UNCHANGED_EXISTING_ROWS=472')
    print('NEW_UNSIGNED_DRAFT_ROWS=13 IDs_UNIQUE=PASS')
    print('UNMODIFIED_SHEETS=',','.join(unaffected))
    print('OOXML_CHANGED_PARTS=',changed)
    print('CELL_DELTA=',len(delta),'NUMERIC_RECHECK_QUEUE=',len(selected))
    print('ORIGINAL_SHA=',sha(B.read_bytes()),'UPDATED_SHA=',sha(X.read_bytes()))
    print('OUTPUTS=',A/'R8A_RBI_FINAL_AUDIT.md',A/'R8A_RBI_CELL_DELTA.csv')
if __name__=='__main__': main()
