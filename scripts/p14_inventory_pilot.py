"""Read-only P14 source inventory and 20-page extraction pilot.

Writes only to work/data_staging/p14 and docs/data_quality/P14_*.csv.
Raw upstream PDF/XLSX and teacher code/data are opened read-only.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import subprocess
import sys
import time
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
UPSTREAM = ROOT / "upstream"
STAGE = ROOT / "work" / "data_staging" / "p14"
DOCS = ROOT / "docs" / "data_quality"
PILOT = [
    # 8 GB/T 30579 pages stratified across chapters 5–9; page-index is 0-based.
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 7, "DM 5.1 / chapter 5"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 32, "DM 5.19 / chapter 5"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 56, "DM 6.7 / chapter 6"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 63, "DM 6.12 / chapter 6"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 71, "DM 7.3 / chapter 7"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 76, "DM 7.8 / chapter 7"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 97, "DM 8.6 / chapter 8"),
    ("02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf", 108, "DM 9.3 / chapter 9"),
    # 8 pages from RBI 2/.4/.5 tables and limits.
    ("01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf", 7, "RBI inspection strategy"),
    ("01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf", 15, "RBI inspection strategy"),
    ("01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf", 29, "RBI likelihood table"),
    ("01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf", 78, "RBI likelihood table"),
    ("01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf", 128, "RBI limits/table"),
    ("01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf", 19, "RBI consequence table"),
    ("01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf", 34, "RBI consequence table"),
    ("01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf", 46, "RBI consequence table"),
    # 4 pages: TSG and supplementary equipment/corrosion standards.
    ("01_RBI标准/补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf", 80, "TSG periodic inspection"),
    ("02_损伤机理参考/补充标准/GBT150.1-2024_压力容器第1部分通用要求.pdf", 16, "supplementary applicability"),
    ("02_损伤机理参考/补充标准/GBT151-2026_热交换器.pdf", 80, "supplementary applicability"),
    ("02_损伤机理参考/补充标准/GBT21447-2018_钢质管道外腐蚀控制规范.pdf", 8, "supplementary corrosion control"),
]


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")


def _csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def _native_text_quality(text: str) -> tuple[str, float, int]:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return "empty", 0.0, 0
    suspicious = sum(
        c == "\ufffd" or "\ue000" <= c <= "\uf8ff" or (not c.isprintable())
        for c in chars
    )
    ratio = suspicious / len(chars)
    if len(chars) < 80:
        label = "short_or_visual"
    elif ratio > 0.02:
        label = "encoding_suspect"
    else:
        # Presence and printability do not prove semantic correctness: embedded CMaps
        # may return readable but wrong glyphs. Every page remains unreviewed.
        label = "native_text_present_unverified"
    return label, round(ratio, 4), len(chars)


def _bottom_page_number(page: fitz.Page) -> int | None:
    """Return only a bottom-line numeric candidate, never claim it is verified."""
    h = page.rect.height
    candidates: list[int] = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        x0, y0, x1, y1 = block.get("bbox", (0, 0, 0, 0))
        if y0 < h * 0.88:
            continue
        s = "".join(span.get("text", "") for line in block.get("lines", []) for span in line.get("spans", [])).strip()
        if re.fullmatch(r"[0-9０-９]{1,3}", s):
            candidates.extend(int(c.translate(str.maketrans("０１２３４５６７８９", "0123456789"))) for c in s)
    return candidates[0] if len(candidates) == 1 else None


def _headings(text: str) -> list[str]:
    found: list[str] = []
    for line in text.splitlines():
        s = " ".join(line.split())
        if re.match(r"^(?:第\s*[一二三四五六七八九十0-9]+\s*章|[1-9](?:\.\d+){0,3}\s+\S)", s):
            if len(s) <= 120:
                found.append(s)
    return found[:8]


def _table_markers(text: str) -> list[str]:
    return list(dict.fromkeys(re.findall(r"(?:表|图|公式|式)\s*[A-ZＡ-Ｚ]?[0-9０-９]+(?:[.-][0-9０-９]+)*", text)))[:40]


def _source_pdfs() -> list[Path]:
    return sorted([* (UPSTREAM / "01_RBI标准").rglob("*.pdf"), * (UPSTREAM / "02_损伤机理参考").rglob("*.pdf")])


def run_inventory() -> None:
    STAGE.mkdir(parents=True, exist_ok=True)
    pdfs = _source_pdfs()
    expected = 12
    inventory: list[dict[str, Any]] = []
    page_rows: list[dict[str, Any]] = []
    cache_path = STAGE / "page_cache" / "native_pages.jsonl"
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with cache_path.open("w", encoding="utf-8", newline="\n") as cache:
        for path in pdfs:
            rel = path.relative_to(UPSTREAM).as_posix()
            doc = fitz.open(path)
            raw = path.read_bytes()
            lens: list[int] = []
            visual_pages = 0
            per_pdf_pages: list[dict[str, Any]] = []
            for idx, page in enumerate(doc):
                text = page.get_text("text", sort=True)
                quality, ratio, count = _native_text_quality(text)
                lens.append(count)
                if quality in {"empty", "short_or_visual", "encoding_suspect", "layout_or_encoding_review"}:
                    visual_pages += 1
                headings = _headings(text)
                markers = _table_markers(text)
                printed_guess = _bottom_page_number(page)
                blocks = []
                for block in page.get_text("dict").get("blocks", []):
                    if block.get("type") == 0:
                        btext = "".join(span.get("text", "") for line in block.get("lines", []) for span in line.get("spans", []))
                        if btext.strip():
                            blocks.append({"bbox": [round(float(x), 2) for x in block.get("bbox", [])], "text": btext[:2000]})
                record = {
                    "source_pdf": rel, "pdf_sha256": sha256(raw), "pdf_page_index": idx,
                    "page_width_pt": round(page.rect.width, 2), "page_height_pt": round(page.rect.height, 2),
                    "printed_page_guess": printed_guess, "printed_page_status": "unverified_guess" if printed_guess is not None else "unknown",
                    "native_text_quality": quality, "native_text_cjk_printability_ratio": ratio,
                    "native_text_char_count": count, "headings_machine_detected": headings,
                    "table_figure_markers_machine_detected": markers, "evidence_level": "machine_located",
                    "human_visual_review": False, "text": text, "blocks": blocks,
                }
                cache.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
                row = {k: record[k] for k in ("source_pdf", "pdf_sha256", "pdf_page_index", "printed_page_guess", "printed_page_status", "native_text_quality", "native_text_cjk_printability_ratio", "native_text_char_count", "headings_machine_detected", "table_figure_markers_machine_detected", "evidence_level", "human_visual_review")}
                row["candidate_status"] = "no_claim_candidate_extracted"
                row["review_status"] = "pending_human_or_visual_verification"
                row["review_reason"] = "native extraction is not evidence of exact quote; printed page and table/cell mapping unverified"
                per_pdf_pages.append(row)
            doc.close()
            pages = len(lens)
            inventory.append({"source_pdf": rel, "sha256": sha256(raw), "bytes": len(raw), "pdf_pages": pages,
                              "pages_text_ge_80_chars": sum(n >= 80 for n in lens), "pages_text_lt_80_chars": sum(n < 80 for n in lens),
                              "total_native_text_chars": sum(lens), "pages_flagged_visual_or_encoding_review": visual_pages,
                              "scope_status": "in_scope_read_only", "human_reviewed_pages": 0})
            page_rows.extend(per_pdf_pages)
    _csv(STAGE / "source_inventory.csv", inventory, list(inventory[0]) if inventory else [])
    _csv(DOCS / "P14_SOURCE_COVERAGE_MATRIX.csv", page_rows, list(page_rows[0]) if page_rows else [])
    # Preserve the exact teacher templates in staging. RBI uses the tracked HEAD baseline (472 rows),
    # not the currently modified upstream worktree copy; this does not touch that worktree.
    rbi_rel = "01_RBI标准/条款摘录表_RBI.xlsx"
    rbi_bytes = subprocess.check_output(["git", "-C", str(UPSTREAM), "show", f"HEAD:{rbi_rel}"])
    dm_source = UPSTREAM / "02_损伤机理参考" / "条款摘录表_损伤机理.xlsx"
    dm_bytes = dm_source.read_bytes()
    rbi_source_copy = STAGE / "source_baselines" / "RBI_HEAD_472.xlsx"
    dm_source_copy = STAGE / "source_baselines" / "DM_source.xlsx"
    rbi_source_copy.parent.mkdir(parents=True, exist_ok=True)
    rbi_source_copy.write_bytes(rbi_bytes)
    dm_source_copy.write_bytes(dm_bytes)
    (STAGE / "条款摘录表_RBI_candidate.xlsx").write_bytes(rbi_bytes)
    (STAGE / "条款摘录表_损伤机理_candidate.xlsx").write_bytes(dm_bytes)
    templates = []
    for name in ("条款摘录表_RBI_candidate.xlsx", "条款摘录表_损伤机理_candidate.xlsx"):
        book = load_workbook(STAGE / name, read_only=True, data_only=False)
        sheet = book["摘录表"]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        records = sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if row and row[0])
        templates.append({"candidate_workbook": name, "sheets": "|".join(book.sheetnames), "sheet_count": len(book.sheetnames), "extract_sheet_columns": len(header), "extract_rows": records, "header_sha256": sha256(json.dumps(header, ensure_ascii=False, default=str).encode())})
        book.close()
    _csv(STAGE / "template_check.csv", templates, list(templates[0]))
    if len(pdfs) != expected:
        raise RuntimeError(f"expected {expected} PDFs, found {len(pdfs)}")
    print(f"PDFs={len(pdfs)} pages={sum(x['pdf_pages'] for x in inventory)} page_cache={cache_path}")
    print(f"matrix={DOCS / 'P14_SOURCE_COVERAGE_MATRIX.csv'} rbi_head_sha256={sha256(rbi_bytes)}")
    for row in inventory:
        print(f"{row['source_pdf']}\t{row['pdf_pages']} pages\tvisual/encoding flags={row['pages_flagged_visual_or_encoding_review']}")


def run_pilot(render_scale: float = 1.15) -> None:
    if not (STAGE / "page_cache" / "native_pages.jsonl").exists():
        raise RuntimeError("run inventory first")
    pilot_dir = STAGE / "pilot"
    image_dir = pilot_dir / "rendered_pages"
    image_dir.mkdir(parents=True, exist_ok=True)
    by_source: dict[str, fitz.Document] = {}
    selected = []
    for rel, idx, purpose in PILOT:
        path = UPSTREAM / Path(rel)
        by_source.setdefault(rel, fitz.open(path))
        doc = by_source[rel]
        if idx >= len(doc):
            raise RuntimeError(f"pilot index out of range: {rel}:{idx}")
        page = doc[idx]
        filename = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(rel).stem) + f"__p{idx:03d}.png"
        target = image_dir / filename
        pix = page.get_pixmap(matrix=fitz.Matrix(render_scale, render_scale), alpha=False)
        pix.save(target)
        text = page.get_text("text", sort=True)
        quality, ratio, count = _native_text_quality(text)
        selected.append({"source_pdf": rel, "pdf_page_index": idx, "printed_page_guess": _bottom_page_number(page), "printed_page_status": "unverified_guess", "purpose": purpose,
                         "native_text_quality": quality, "native_text_char_count": count, "native_text_cjk_printability_ratio": ratio,
                         "render_path": target.relative_to(ROOT).as_posix(), "visual_review_status": "pending_manual_review",
                         "candidate_count": 0, "precision": "not_scored_no_claims_extracted"})
    for doc in by_source.values():
        doc.close()
    _csv(pilot_dir / "pilot_20_pages.csv", selected, list(selected[0]))
    # Contact sheets are only navigation aids; original page renders remain in the private staging folder.
    try:
        from PIL import Image, ImageDraw
        thumbs = []
        for row in selected:
            im = Image.open(ROOT / row["render_path"]).convert("RGB")
            im.thumbnail((250, 350))
            canvas = Image.new("RGB", (270, 390), "white")
            canvas.paste(im, ((270-im.width)//2, 22))
            ImageDraw.Draw(canvas).text((6, 4), f"idx {row['pdf_page_index']} | {Path(row['source_pdf']).name[:22]}", fill="black")
            thumbs.append(canvas)
        sheet = Image.new("RGB", (270*4, 390*5), "#d8dedb")
        for i, im in enumerate(thumbs):
            sheet.paste(im, ((i%4)*270, (i//4)*390))
        sheet.save(pilot_dir / "pilot_contact_sheet.png")
    except ImportError:
        pass
    print(f"pilot_pages={len(selected)} rendered={image_dir} contact={pilot_dir / 'pilot_contact_sheet.png'}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["inventory", "pilot"])
    parser.add_argument("--render-scale", type=float, default=1.15)
    args = parser.parse_args()
    t0 = time.perf_counter()
    if args.mode == "inventory":
        run_inventory()
    else:
        run_pilot(args.render_scale)
    print(f"elapsed_s={time.perf_counter()-t0:.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
