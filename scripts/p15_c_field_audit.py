"""Read-only P15 field QA; never mutates upstream workbooks or official JSONL."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
UP = ROOT / "upstream"
RBI = UP / "01_RBI标准" / "条款摘录表_RBI.xlsx"
DM = UP / "02_损伤机理参考" / "条款摘录表_损伤机理.xlsx"
CLAUSES = UP / "05_提交区" / "clauses.jsonl"
ENUMS = json.loads((UP / "04_接口文档" / "enums.json").read_text(encoding="utf-8"))
EXPECTED_SHEETS = ["填写说明", "摘录表", "填写示例", "枚举值"]
RBI_HEADERS = ["clause_id", "standard_code", "clause_no", "page", "title", "原文摘录", "失效可能性等级", "失效后果等级", "风险等级(原始)", "风险等级(统一)", "检验周期(年)", "推荐检验方法", "适用材质", "适用介质", "温度下限", "温度上限", "压力下限", "压力上限", "适用设备类型", "其他适用条件", "抽取人", "复核人", "备注"]
DM_HEADERS = ["clause_id", "standard_code", "clause_no", "page", "title", "原文摘录", "损伤描述及机理", "损伤形态", "受影响的材料", "主要影响因素", "易发生部位", "主要预防措施", "检测监测方法", "相关或伴随损伤", "适用材质", "适用介质", "温度下限", "温度上限", "压力下限", "压力上限", "适用设备类型", "其他适用条件", "抽取人", "复核人", "备注"]
RBI_POLICY = {"clause_id":"REQUIRED", "standard_code":"REQUIRED", "clause_no":"REQUIRED", "page":"REQUIRED", "title":"REQUIRED", "原文摘录":"REQUIRED", "失效可能性等级":"CONDITIONAL_SUPPORTED", "失效后果等级":"CONDITIONAL_SUPPORTED", "风险等级(原始)":"CONDITIONAL_SUPPORTED", "风险等级(统一)":"CONDITIONAL_SUPPORTED", "检验周期(年)":"CONDITIONAL_SUPPORTED", "推荐检验方法":"CONDITIONAL_SUPPORTED", "适用材质":"CONDITIONAL_SUPPORTED", "适用介质":"CONDITIONAL_SUPPORTED", "温度下限":"CONDITIONAL_SUPPORTED", "温度上限":"CONDITIONAL_SUPPORTED", "压力下限":"CONDITIONAL_SUPPORTED", "压力上限":"CONDITIONAL_SUPPORTED", "适用设备类型":"CONDITIONAL_SUPPORTED", "其他适用条件":"CONDITIONAL_SUPPORTED", "抽取人":"REQUIRED", "复核人":"REVIEW_NEEDED", "备注":"CONDITIONAL_SUPPORTED"}
DM_POLICY = {"clause_id":"REQUIRED", "standard_code":"REQUIRED", "clause_no":"REQUIRED", "page":"REQUIRED", "title":"REQUIRED", "原文摘录":"REQUIRED", **{x:"REQUIRED" for x in DM_HEADERS[6:14]}, "适用材质":"CONDITIONAL_SUPPORTED", "适用介质":"CONDITIONAL_SUPPORTED", "温度下限":"CONDITIONAL_SUPPORTED", "温度上限":"CONDITIONAL_SUPPORTED", "压力下限":"CONDITIONAL_SUPPORTED", "压力上限":"CONDITIONAL_SUPPORTED", "适用设备类型":"CONDITIONAL_SUPPORTED", "其他适用条件":"CONDITIONAL_SUPPORTED", "抽取人":"REQUIRED", "复核人":"REVIEW_NEEDED", "备注":"CONDITIONAL_SUPPORTED"}
PDF_BY_STANDARD = {
 "GB/T 26610.1-2022":"01_RBI标准/GBT26610.1-2022_基本要求和实施程序.pdf", "GB/T 26610.2-2022":"01_RBI标准/GBT26610.2-2022_基于风险的检验策略.pdf", "GB/T 26610.3-2014":"01_RBI标准/GBT26610.3-2014_风险的定性分析方法.pdf", "GB/T 26610.4-2022":"01_RBI标准/GBT26610.4-2022_失效可能性定量分析方法.pdf", "GB/T 26610.5-2022":"01_RBI标准/GBT26610.5-2022_失效后果定量分析方法.pdf", "TSG 21-2016":"01_RBI标准/补充规范/TSG21-2016_固定式压力容器安全技术监察规程.pdf", "GB/T 30579-2022":"02_损伤机理参考/GBT30579-2022_承压设备损伤模式识别.pdf"
}

def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""): h.update(block)
    return h.hexdigest()

def txt(v: Any) -> str:
    return "" if v is None else str(v).strip()

def iter_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if wb.sheetnames != EXPECTED_SHEETS:
        yield "__SHEET_ERROR__", {"expected": EXPECTED_SHEETS, "actual": wb.sheetnames}; return
    ws = wb["摘录表"]
    rows = ws.iter_rows(values_only=True)
    headers = [txt(v) for v in next(rows)]
    for excel_row, values in enumerate(rows, start=2):
        row = dict(zip(headers, values))
        if any(txt(v) for v in row.values()): yield excel_row, row

def load_risk_map() -> dict:
    # Independent read-only parse of the teacher mapping workbook (avoid executing/importing teacher code).
    path=UP/"04_接口文档"/"risk_level_mapping.xlsx"
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb["GBT26610.3-2014"]; rows=ws.iter_rows(values_only=True); next(rows,None)
        return {(str(a).strip(),str(b).strip()):(str(c).strip(),str(d).strip()) for a,b,c,d,*_ in rows if a is not None and b is not None}
    finally: wb.close()

def classify_nonempty(kind: str, field: str, value: Any, row: dict) -> tuple[str, str]:
    v = txt(value)
    if field in ("复核人",): return ("REVIEW_NEEDED", "保持未签；需真实独立复核人")
    if not v:
        policy = RBI_POLICY if kind == "rbi" else DM_POLICY
        if policy.get(field) == "REQUIRED": return ("REVIEW_NEEDED", "必需字段为空")
        return ("SOURCE_UNSPECIFIED", "空值保留；需来源支持后才填写")
    if field in ("温度下限", "温度上限", "压力下限", "压力上限", "检验周期(年)"):
        return ("REVIEW_NEEDED", "数值/单位/变量对象需逐项对照 PDF；表中数字不自动证明语义")
    if field == "其他适用条件" and re.search(r"拆行|顺序|等级划分见表|表[0-9一二三四五六七八九]+中", v):
        return ("REVIEW_NEEDED", "疑似表格解释/管理备注混入适用条件，须核原文/条件对象")
    if field == "原文摘录": return ("REVIEW_NEEDED", "机器只能做长度/定位检查，不能证明逐字忠实")
    if field in ("适用材质", "适用介质", "适用设备类型"):
        table = {"适用材质":"material", "适用介质":"medium", "适用设备类型":"equipment_type"}[field]
        legal = set(ENUMS[table]); aliases={alias:key for key,value in ENUMS[table].items() for alias in value.get("aliases",[])}
        vals = [x.strip() for x in re.split(r"[;；]", v) if x.strip()]
        invalid = [x for x in vals if x not in legal and x not in aliases]
        if invalid: return ("REVIEW_NEEDED", "枚举非法: " + ";".join(invalid))
        noncanonical=[f"{x}->{aliases[x]}" for x in vals if x in aliases]
        if noncanonical: return ("REVIEW_NEEDED", "枚举别名需验证教师归一化并确认语义: "+";".join(noncanonical))
    return ("CONDITIONAL_SUPPORTED", "值存在；来源定位/适用性/方向语义仍需独立证据")

def audit_workbook(path: Path, kind: str, out: Path) -> tuple[int, list[dict]]:
    expected = RBI_HEADERS if kind == "rbi" else DM_HEADERS
    policy = RBI_POLICY if kind == "rbi" else DM_POLICY
    ledger=[]; seen=set(); ids=[]
    wb=load_workbook(path,read_only=True,data_only=True)
    if wb.sheetnames != EXPECTED_SHEETS: ledger.append({"severity":"ERROR","file":str(path),"field":"sheets","message":f"sheet mismatch {wb.sheetnames}"})
    ws=wb["摘录表"]; it=ws.iter_rows(values_only=True); headers=[txt(v) for v in next(it)]
    if headers != expected: ledger.append({"severity":"ERROR","file":str(path),"field":"header","message":"header/order mismatch"})
    for rownum, values in enumerate(it,start=2):
        row=dict(zip(headers,values))
        if not any(txt(v) for v in row.values()): continue
        cid=txt(row.get("clause_id")); ids.append(cid)
        if cid in seen: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"clause_id","message":"duplicate primary key"})
        seen.add(cid)
        if len(txt(row.get("原文摘录")))>300: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"原文摘录","message":"F exceeds 300 chars"})
        page=row.get("page")
        if not isinstance(page,(int,float)) or int(page)<1: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"page","message":f"invalid printed page {page!r}"})
        if kind=="rbi":
            g,h=txt(row.get("失效可能性等级")),txt(row.get("失效后果等级"))
            if g and h and txt(row.get("standard_code"))=="GB/T 26610.3-2014":
                expected_risk=load_risk_map().get((g,h))
                if not expected_risk: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"G:J","message":f"matrix cell ({g},{h}) absent from teacher mapping"})
                elif txt(row.get("风险等级(原始)")) and txt(row.get("风险等级(原始)"))!=expected_risk[0]: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"风险等级(原始)","message":f"matrix says {expected_risk[0]}, workbook says {row.get('风险等级(原始)')}"})
                elif txt(row.get("风险等级(统一)")) and txt(row.get("风险等级(统一)"))!=expected_risk[1]: ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"风险等级(统一)","message":f"matrix says {expected_risk[1]}, workbook says {row.get('风险等级(统一)')}"})
            if txt(row.get("standard_code")) in ("GB/T 26610.4-2022",) and any(txt(row.get(x)) for x in ("失效可能性等级","失效后果等级","风险等级(原始)","风险等级(统一)")):
                ledger.append({"severity":"REVIEW","file":str(path),"row":rownum,"clause_id":cid,"field":"G:J","message":".4定量概率/因子不得未经语义证明映射为等级/矩阵风险"})
        else:
            temp_lo=row.get("温度下限"); temp_hi=row.get("温度上限")
            if (temp_lo is not None and temp_hi is not None) and float(temp_lo)>float(temp_hi): ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"温度下限/上限","message":"lower bound exceeds upper bound"})
            if (row.get("压力下限") is not None and row.get("压力上限") is not None) and float(row["压力下限"])>float(row["压力上限"]): ledger.append({"severity":"ERROR","file":str(path),"row":rownum,"clause_id":cid,"field":"压力下限/上限","message":"lower bound exceeds upper bound"})
            if cid.endswith("-5.23") or cid.endswith("-5.6") or cid.endswith("-6.7"):
                ledger.append({"severity":"REVIEW","file":str(path),"row":rownum,"clause_id":cid,"field":"temperature/pressure/material/other","message":"priority semantic review: local phenomenon vs general range; H2S partial vs total pressure; susceptibility vs resistant material"})
        for field in expected:
            value=row.get(field)
            cls,reason=classify_nonempty(kind,field,value,row)
            if cls=="REVIEW_NEEDED": ledger.append({"severity":"REVIEW","file":str(path),"row":rownum,"clause_id":cid,"field":field,"message":reason})
            pdf_rel=PDF_BY_STANDARD.get(txt(row.get("standard_code")),"")
            pdf_path=UP/pdf_rel if pdf_rel else None
            pdf_exists=bool(pdf_path and pdf_path.is_file())
            ledger.append({"severity":"CELL","kind":kind,"file":path.name,"row":rownum,"clause_id":cid,"field":field,"value":txt(value),"field_class":cls,"policy":policy.get(field,"REVIEW_NEEDED"),"source_pdf":pdf_rel,"pdf_exists":pdf_exists,"printed_page":txt(row.get("page")),"pdf_page_index":"","bbox_or_table_cell":"","evidence_quote":txt(row.get("原文摘录")),"review_status":"PENDING_HUMAN_VISUAL","basis":"cell-level policy; PDF file presence/page number only is not proof of locator or source fidelity"})
    if len(ids)!=len(set(ids)): pass
    out.parent.mkdir(parents=True,exist_ok=True)
    with out.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=sorted({k for x in ledger for k in x})); w.writeheader(); w.writerows(ledger)
    return sum(1 for x in ledger if x["severity"]=="CELL"),ledger

def audit_official_clause_semantics(out:Path):
    records=[json.loads(line) for line in CLAUSES.read_text(encoding="utf-8").splitlines() if line.strip()]
    by_id={r["clause_id"]:r for r in records}; issues=[]; riskmap=load_risk_map()
    matrix=[r for r in records if r.get("clause_id","").startswith("GBT26610.3-2014-7-")]
    matrix_expected=[(str(l),c) for l in range(1,6) for c in "ABCDE"]
    seen=[]
    for r in matrix:
        detail=r.get("rbi_detail") or {}; pair=(txt(detail.get("likelihood")),txt(detail.get("consequence"))); seen.append(pair)
        expected=riskmap.get(pair)
        if expected is None or (detail.get("risk_level"),detail.get("risk_level_unified"))!=expected:
            issues.append({"severity":"ERROR","rule":"RISK_MATRIX_MAPPING","clause_id":r["clause_id"],"observed":json.dumps(detail,ensure_ascii=False),"expected":str(expected)})
    if sorted(seen)!=sorted(matrix_expected): issues.append({"severity":"ERROR","rule":"RISK_MATRIX_5X5_COMPLETENESS","observed":str(sorted(seen)),"expected":str(matrix_expected)})
    area=[r for r in records if r.get("clause_id","").startswith("GBT26610.5-2022-表2-")]
    expected_area={"a":"CA≤9.29","b":"9.29<CA≤92.9","c":"92.9<CA≤279","d":"279<CA≤929","e":"CA>929"}
    for suffix,fragment in expected_area.items():
        cid=f"GBT26610.5-2022-表2-{suffix}"; row=by_id.get(cid)
        if not row or fragment not in row.get("content_raw",""):
            issues.append({"severity":"ERROR","rule":"AREA_CONSEQUENCE_THRESHOLD","clause_id":cid,"observed":row.get("content_raw","") if row else "MISSING","expected_fragment":fragment})
    dm_checks={"GBT30579-2022-5.23":{"temperature_c":[49,66],"semantic_review":"49-66 C may be NH4HS gas-phase precipitation sub-phenomenon, not blanket mechanism boundary"},"GBT30579-2022-5.6":{"temperature_c":[177,242],"other_contains":"218℃～400℃","semantic_review":"general range vs light-phase condition object needs independent PDF adjudication"},"GBT30579-2022-6.7":{"pressure_mpa":[None,None],"other_contains":"0.0003 MPa","semantic_review":"0.0003 MPa is recorded as H2S partial pressure; must not be total operating pressure bound"}}
    for cid,check in dm_checks.items():
        row=by_id.get(cid); cond=(row or {}).get("applicable_conditions",{})
        passed=bool(row)
        if "temperature_c" in check: passed &= cond.get("temperature_c")==check["temperature_c"]
        if "pressure_mpa" in check: passed &= cond.get("pressure_mpa")==check["pressure_mpa"]
        if "other_contains" in check: passed &= check["other_contains"] in str(cond.get("other",""))
        issues.append({"severity":"REVIEW" if passed else "ERROR","rule":"DM_HIGH_RISK_SEMANTICS","clause_id":cid,"observed":json.dumps(cond,ensure_ascii=False),"expected":check["semantic_review"],"structural_assertion_pass":passed,"source_truth":"NOT_INDEPENDENTLY_VISUAL_VERIFIED"})
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps({"official_clause_count":len(records),"rbi_matrix_cells":len(matrix),"matrix_exact_5x5_pairs":len(set(seen))==25 and sorted(seen)==sorted(matrix_expected),"area_table_rows":len(area),"issues":issues,"interpretation":"Numeric consistency checks against current teacher data and mapping only; not evidence that PDFs were faithfully transcribed."},ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    errors=sum(x["severity"]=="ERROR" for x in issues); reviews=sum(x["severity"]=="REVIEW" for x in issues)
    print(f"official_semantic_audit clauses={len(records)} matrix_cells={len(matrix)} area_rows={len(area)} machine_errors={errors} high_risk_review={reviews} report={out}")
    return issues

def main(argv=None):
    p=argparse.ArgumentParser(); p.add_argument("--out",type=Path,default=ROOT/"work/data_staging/p15/C_QA/P15_C_FIELD_LEDGER.csv"); p.add_argument("--rbi-input",type=Path,default=RBI); p.add_argument("--dm-input",type=Path,default=DM); a=p.parse_args(argv)
    total=0; all_issues=[]
    tag="candidate_" if a.rbi_input.resolve()!=RBI.resolve() or a.dm_input.resolve()!=DM.resolve() else ""
    for path,kind in ((a.rbi_input,"rbi"),(a.dm_input,"dm")):
        n,issues=audit_workbook(path,kind,a.out.with_name(f"{tag}{kind}_"+a.out.name)); total+=n; all_issues+=issues
        print(f"{kind}: source_sha256={sha(path)} cell_rows={n} errors={sum(i['severity']=='ERROR' for i in issues)} review={sum(i['severity']=='REVIEW' for i in issues)}")
    if not tag: audit_official_clause_semantics(a.out.parent/"P15_C_OFFICIAL_SEMANTIC_CHECKS.json")
    print(f"cell_policy_rows={total}; field policy labels do not assert source truth; outputs={a.out.parent}")
    return 0
if __name__=="__main__": raise SystemExit(main())
