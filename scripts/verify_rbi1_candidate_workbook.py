"""Read-only audit of the isolated 13-row RBI1 workbook against its PDF.

This checks workbook/template integrity and text on the page; visual title and
clause-context review is recorded separately in the P2 evidence document.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path

import fitz
from openpyxl import load_workbook

from scripts.audit_rbi1_quote_candidates import CANDIDATES, PDF

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "upstream" / "01_RBI标准" / "条款摘录表_RBI.xlsx"
WORKING = ROOT / "upstream" / "01_RBI标准" / "RBI_GBT26610_1_13candidates_for_review.xlsx"


def key(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text or ""))


def run() -> int:
    src_wb = load_workbook(SOURCE, read_only=True, data_only=True)
    work_wb = load_workbook(WORKING, read_only=True, data_only=True)
    src_ws, ws = src_wb["摘录表"], work_wb["摘录表"]
    src_rows = list(src_ws.iter_rows(values_only=True))
    rows = list(ws.iter_rows(values_only=True))
    headers = tuple(str(v or "") for v in rows[0])
    if headers != tuple(str(v or "") for v in src_rows[0]):
        raise AssertionError("candidate headers/order differ from the original 23-column RBI sheet")
    records = [row for row in rows[1:] if any(v not in (None, "") for v in row)]
    if len(records) != len(CANDIDATES):
        raise AssertionError(f"expected {len(CANDIDATES)} candidate rows, got {len(records)}")

    known_ids = {str(row[0]) for row in src_rows[1:] if row and row[0] not in (None, "")}
    candidate_ids: list[str] = []
    pdf = fitz.open(PDF)
    for row, (printed_page, clause, title, quote) in zip(records, CANDIDATES, strict=True):
        clause_id, standard, row_clause, page, row_title, raw = row[:6]
        if (str(standard), str(row_clause), int(page), str(row_title), str(raw)) != (
            "GB/T 26610.1-2022", clause, printed_page, title, quote
        ):
            raise AssertionError(f"candidate row mismatch for {clause}")
        pdf_index = printed_page + 4
        page_text = pdf[pdf_index].get_text(sort=True)
        if key(quote) not in key(page_text):
            raise AssertionError(f"quote mismatch: {clause} page {printed_page}")
        if row[20] not in (None, "") or row[21] not in (None, ""):
            raise AssertionError(f"extractor/reviewer identity must stay blank: {clause}")
        if any(row[i] not in (None, "") for i in range(6, 19)):
            raise AssertionError(f"unsupported risk/condition value must stay blank: {clause}")
        if not isinstance(row[2], str):
            raise AssertionError(f"clause_no must be text: {clause}")
        cid = str(clause_id)
        if cid in known_ids:
            raise AssertionError(f"candidate duplicates an ID already in original workbook: {cid}")
        candidate_ids.append(cid)
        print(f"MACHINE_MATCH {clause} printed_page={printed_page} pdf_index={pdf_index} quote=PASS title=VISUAL_REVIEW_REQUIRED")
    if len(candidate_ids) != len(set(candidate_ids)):
        raise AssertionError("duplicate candidate clause_id")

    for wb in (src_wb, work_wb):
        wb.close()
    print("MACHINE_DRAFT_CHECK_PASS rows=13 headers=23 IDs=unique no_source_overlap identities=blank risk_fields=blank")
    print("SOURCE_SHA256", hashlib.sha256(SOURCE.read_bytes()).hexdigest())
    print("WORKING_COPY_SHA256", hashlib.sha256(WORKING.read_bytes()).hexdigest())
    print("PDF_SHA256", hashlib.sha256(PDF.read_bytes()).hexdigest())
    return 0


if __name__ == "__main__":
    raise SystemExit(run())
