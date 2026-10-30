"""Build a page-complete OCR/structure index and clause-to-XLSX audit ledgers.

OCR is discovery evidence only: it never certifies wording, numeric values,
printed pagination, or human review. Unclear pages stay explicitly unresolved.
The script does not modify the teacher workbook or official JSONL.
"""
from __future__ import annotations

import csv
import hashlib
from io import BytesIO
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
UP = ROOT / "upstream"
RBI = UP / "01_RBI标准"
AUDIT = ROOT / "docs" / "audits"
PDFS = [
    ("GB/T 26610.1-2022", RBI / "GBT26610.1-2022_基本要求和实施程序.pdf"),
    ("GB/T 26610.2-2022", RBI / "GBT26610.2-2022_基于风险的检验策略.pdf"),
    ("GB/T 26610.3-2014", RBI / "GBT26610.3-2014_风险的定性分析方法.pdf"),
    ("GB/T 26610.4-2022", RBI / "GBT26610.4-2022_失效可能性定量分析方法.pdf"),
    ("GB/T 26610.5-2022", RBI / "GBT26610.5-2022_失效后果定量分析方法.pdf"),
    ("TSG 21-2016", RBI / "补充规范" / "TSG21-2016_固定式压力容器安全技术监察规程.pdf"),
]
XLSX = RBI / "条款摘录表_RBI.xlsx"
LEDGER = AUDIT / "R8A_RBI_PDF_COVERAGE_LEDGER.csv"
CELL_DELTA = AUDIT / "R8A_RBI_CELL_DELTA.csv"
SUMMARY = AUDIT / "R8A_RBI_PAGE_INVENTORY.csv"
OUT_MD = AUDIT / "R8A_RBI_FINAL_AUDIT.md"
COLS = list("ABCDEFGHIJKLMNOPQRSTUVW")


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def norm(s: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s)).lower()


def ocr_page(page: fitz.Page, temp: Path) -> str:
    pix = page.get_pixmap(matrix=fitz.Matrix(1.55, 1.55), alpha=False)
    image = temp / "page.png"
    pix.save(image)
    proc = subprocess.run(
        ["tesseract", str(image), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", check=False,
    )
    image.unlink(missing_ok=True)
    if proc.returncode:
        raise RuntimeError(f"tesseract failed ({proc.returncode}): {proc.stderr[-500:]}")
    return proc.stdout


def printed_page(text: str) -> str:
    # Do not infer an offset. Only accept explicit footer/page-marker patterns.
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    for line in reversed(lines[-8:]):
        m = re.search(r"(?:—|-|－)\s*(\d{1,3})\s*(?:—|-|－)", line)
        if m:
            return m.group(1)
        m = re.fullmatch(r"(?:第\s*)?(\d{1,3})(?:\s*页)?", line)
        if m:
            return m.group(1)
    return ""


def section_items(text: str) -> list[str]:
    out = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            continue
        if re.match(r"^(?:\d{1,2}(?:\.\d{1,3}){0,4}|[A-Z]\s*\.\s*\d+(?:\.\d+)*)\s+\S", line):
            out.append(line[:180])
        elif re.match(r"^(?:表|图)\s*[A-Z]?\s*\d+(?:[.．]\d+)*", line):
            out.append(line[:180])
        elif re.match(r"^(?:公式|式)\s*[（(]?\s*\d+", line):
            out.append(line[:180])
    return list(dict.fromkeys(out))


def main() -> None:
    AUDIT.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["摘录表"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if values and values[0]:
            rec = dict(zip(headers, values))
            rec["_row"] = row_no
            rows.append(rec)
    wb.close()

    by_standard: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_standard[str(row["standard_code"])].append(row)
    candidate_map: dict[tuple[int, str], tuple[str, str]] = {}
    from scripts.audit_rbi1_quote_candidates import CANDIDATES
    for page, clause, title, quote in CANDIDATES:
        candidate_map[(page, clause)] = (title, quote)

    page_records = []
    clause_records = []
    stats = []
    for standard, pdf in PDFS:
        raw_pdf = pdf.read_bytes()
        digest = sha(raw_pdf)
        doc = fitz.open(pdf)
        page_cache = []
        with tempfile.TemporaryDirectory(prefix="r8a_ocr_") as td:
            tmp = Path(td)
            for idx, page in enumerate(doc):
                extracted = page.get_text(sort=True)
                # OCR every page, including garbled-font pages; do not treat OCR as verified text.
                recognized = ocr_page(page, tmp)
                blob = recognized if len(norm(recognized)) > len(norm(extracted)) else extracted
                items = section_items(blob)
                pp = printed_page(recognized) or printed_page(extracted)
                low_text = len(norm(extracted)) < 40
                table_marks = sorted(set(re.findall(r"(?:表|图)\s*[A-Z]?\s*\d+(?:[.．]\d+)*", blob)))
                formula_marks = sorted(set(re.findall(r"(?:式|公式)\s*[（(]?\s*\d+", blob)))
                status = "needs_visual_review"
                if any(x in recognized for x in ("目 次", "目录", "参考文献", "前 言", "引 言")):
                    status = "needs_human_review"
                elif not items and not table_marks and not formula_marks and not low_text:
                    status = "needs_visual_review"
                page_record = {
                    "standard_code": standard, "source_pdf": str(pdf.relative_to(ROOT)),
                    "source_sha256": digest, "pdf_page_0based": idx,
                    "printed_page": pp, "extracted_characters": len(extracted),
                    "ocr_characters": len(recognized), "low_text_layer": str(low_text).lower(),
                    "section_table_formula_markers": " | ".join((items + table_marks + formula_marks)[:80]),
                    "item_count": len(items) + len(table_marks) + len(formula_marks),
                    "review_status": status,
                }
                page_records.append(page_record)
                page_cache.append((idx, pp, extracted, recognized, items, table_marks, formula_marks))

        source_rows = by_standard.get(standard, [])
        used_candidate = set()
        for rec in source_rows:
            # D is the teacher's printed page. It is not accepted as visually verified by this index.
            pno = str(rec.get("page") or "")
            clause = str(rec.get("clause_no") or "")
            matching_pages = [p for p in page_cache if p[1] == pno and clause and norm(clause) in norm(p[2] + "\n" + p[3])]
            exact_quote = any(norm(str(rec.get("原文摘录") or "")) in norm(p[2] + "\n" + p[3]) for p in matching_pages) if rec.get("原文摘录") else False
            clause_records.append({
                "standard_code": standard, "source_pdf": str(pdf.relative_to(ROOT)),
                "source_sha256": digest, "pdf_page_0based": matching_pages[0][0] if matching_pages else "",
                "printed_page": pno, "section_no": clause,
                "table_or_figure": "", "source_item_key": str(rec.get("clause_id")),
                "item_type": "existing_clause", "scope_decision": "existing",
                "existing_clause_id": str(rec.get("clause_id")), "excel_row": rec["_row"],
                "target_field": "A:W", "extraction_method": "existing_xlsx_mapping+page_OCR_diagnostic",
                "source_text_verified": "false", "formula_table_verified": "false",
                "review_status": "needs_human_review", "reason": "现有行无独立复核人；OCR只定位，不证明逐字准确/适用性",
                "notes": "quote_match=" + str(exact_quote).lower() + ";页码映射需视觉核验",
            })

        for (printed, clause), (title, quote) in candidate_map.items():
            if standard != "GB/T 26610.1-2022":
                continue
            p = next((x for x in page_cache if x[1] == str(printed)), None)
            if not p:
                # Known P2 PDF-index mapping; keep this as an explicit unverified candidate page reference.
                p = next((x for x in page_cache if x[0] == printed + 4), None)
            row_no = 474 + len(used_candidate)
            used_candidate.add(clause)
            clause_records.append({
                "standard_code": standard, "source_pdf": str(pdf.relative_to(ROOT)),
                "source_sha256": digest, "pdf_page_0based": p[0] if p else "",
                "printed_page": printed, "section_no": clause,
                "table_or_figure": "", "source_item_key": clause,
                "item_type": "new_candidate", "scope_decision": "extracted",
                "existing_clause_id": "", "excel_row": row_no,
                "target_field": "A-F,W; G-V intentionally blank except no human signatures",
                "extraction_method": "P2_candidate+machine_quote_check+prior_page_visual_review",
                "source_text_verified": "candidate_only", "formula_table_verified": "not_applicable",
                "review_status": "DRAFT_needs_real_extractor_and_independent_review",
                "reason": "无真实抽取/独立复核身份；不作为官方转换输入",
                "notes": f"title={title}; quote_sha256={sha(quote.encode('utf-8'))}; no equipment risk inferred",
            })

        counts = Counter(p[0] for p in page_cache)
        stats.append({
            "standard_code": standard, "source_pdf": str(pdf.relative_to(ROOT)),
            "source_sha256": digest, "page_count": len(doc), "existing_excel_rows": len(source_rows),
            "existing_rows_with_clause_number_OCR_hit": sum(any(p[1] == str(r.get("page")) and str(r.get("clause_no") or "") and norm(str(r.get("clause_no"))) in norm(p[2]+p[3]) for p in page_cache) for r in source_rows),
            "new_candidates": 13 if standard == "GB/T 26610.1-2022" else 0,
            "pages_needing_visual_or_human_review": len(page_cache),
            "pages_with_table_markers": sum(bool(p[5]) for p in page_cache),
            "pages_with_formula_markers": sum(bool(p[6]) for p in page_cache),
            "ocr_pages": len(page_cache),
        })
        doc.close()

    fields = list(clause_records[0])
    with LEDGER.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(clause_records)
    with SUMMARY.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(page_records[0])); w.writeheader(); w.writerows(page_records)
    # The original-to-current delta is generated directly from the unmodified backup.
    backup = XLSX.with_name(XLSX.name + ".bak_owner_before_rbi1_20261028")
    # The safety backup intentionally has a non-XLSX suffix; use in-memory
    # bytes so openpyxl doesn't infer format from the backup filename.
    before_wb = load_workbook(BytesIO(backup.read_bytes()), data_only=False, read_only=True)
    before = before_wb["摘录表"]
    after = load_workbook(XLSX, data_only=False, read_only=True)["摘录表"]
    delta = []
    for r in range(474, 487):
        for col in COLS:
            old = before[f"{col}{r}"].value
            new = after[f"{col}{r}"].value
            if old != new:
                delta.append({"cell": f"{col}{r}", "old": old or "", "new": new or "",
                              "pdf_sha": next(s["source_sha256"] for s in stats if s["standard_code"] == "GB/T 26610.1-2022"),
                              "printed_page": after[f"D{r}"].value or "", "pdf_page": (int(after[f"D{r}"].value) + 4) if after[f"D{r}"].value else "",
                              "confidence_source": "P2 source check; not independent human review",
                              "source_excerpt": "see clause row F; standard text omitted from audit copy",
                              "needs_human_review": "true"})
    before_wb.close(); after.parent.close()
    with CELL_DELTA.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(delta[0])); w.writeheader(); w.writerows(delta)
    print("PDFS", len(stats), "PAGES", sum(x["page_count"] for x in stats))
    print("OCR_PAGE_INDEX", sum(x["ocr_pages"] for x in stats), "CLAUSE_LEDGER_ROWS", len(clause_records), "CELL_DELTAS", len(delta))
    for x in stats: print(x["standard_code"], "pages", x["page_count"], "existing", x["existing_excel_rows"], "candidate", x["new_candidates"], "OCR_clause_hits", x["existing_rows_with_clause_number_OCR_hit"], "visual_review_required", x["pages_needing_visual_or_human_review"])
    print("OUTPUTS", LEDGER, SUMMARY, CELL_DELTA)


if __name__ == "__main__":
    main()
