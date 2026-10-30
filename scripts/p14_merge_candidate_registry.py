"""Merge disjoint machine audit outputs into a P14 L2 registry (never official data)."""
from __future__ import annotations
import csv, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
STAGE=ROOT/'work'/'data_staging'/'p14'
AGENTS=STAGE/'agents'
OUT=STAGE/'candidates'/'p14_source_candidates.jsonl'

def read_jsonl(path):
    with path.open(encoding='utf-8-sig') as f:
        return [json.loads(x) for x in f if x.strip()]
def digest(s): return hashlib.sha256((s or '').encode('utf-8')).hexdigest()
def rows_xlsx(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb['摘录表']; it=ws.iter_rows(values_only=True); header=list(next(it)); values=[]
    for r in it:
        if r and r[0]: values.append(dict(zip(header,r)))
    wb.close(); return values
def normalize_conditions(r, kind):
    if kind=='damage_mechanism':
        return {'applicable_materials':r.get('适用材质'),'medium':r.get('适用介质'),
                'temperature_c':[r.get('温度下限'),r.get('温度上限')],
                'pressure_mpa':[r.get('压力下限'),r.get('压力上限')],
                'equipment_type':r.get('适用设备类型'),'other':r.get('其他适用条件')}
    return {'applicable_materials':r.get('适用材质'),'medium':r.get('适用介质'),
            'temperature_c':[r.get('温度下限'),r.get('温度上限')],
            'pressure_mpa':[r.get('压力下限'),r.get('压力上限')],
            'equipment_type':r.get('适用设备类型'),'other':r.get('其他适用条件'),
            'likelihood':r.get('失效可能性等级'),'consequence':r.get('失效后果等级'),
            'risk_original':r.get('风险等级(原始)'),'risk_unified':r.get('风险等级(统一)'),
            'inspection_interval_year':r.get('检验周期(年)'),'inspection_methods':r.get('推荐检验方法')}

def main():
    dm_rows={r['clause_id']:r for r in rows_xlsx(STAGE/'条款摘录表_损伤机理_candidate.xlsx')}
    # The RBI machine review includes the currently observed 485-row source (472+13 unsigned).
    # Read only; the candidate workbook deliberately uses the clean tracked 472-row HEAD baseline.
    rbi_rows={r['clause_id']:r for r in rows_xlsx(ROOT/'upstream'/'01_RBI标准'/'条款摘录表_RBI.xlsx')}
    dm_screen=read_jsonl(AGENTS/'A'/'dm30579_existing82_machine_candidates.jsonl')
    rbi_audit=list(csv.DictReader((AGENTS/'B'/'rbi_existing_rows_machine_audit.csv').open(encoding='utf-8-sig',newline='')))
    rbi_priority={x['candidate_id']:x for x in read_jsonl(AGENTS/'B'/'rbi_table_quote_review_queue.jsonl')}
    c_rows=read_jsonl(AGENTS/'C'/'conditional_dm_candidates.jsonl')
    out=[]
    for a in dm_screen:
        r=dm_rows.get(a['candidate_id'],{})
        quote=str(r.get('原文摘录') or '')
        out.append({'candidate_id':f"DM-existing-{a['candidate_id']}",'standard_code':a['standard_code'],'clause_no':a['clause_no'],
          'printed_page':a.get('printed_page'),'pdf_page_index':a.get('pdf_page_index'),'source_bbox':a.get('source_bbox'),
          'raw_quote':quote,'raw_quote_sha256':digest(quote),'quote_kind':'text','table_id':None,'row_id':None,'col_id':None,
          'target_sheet':'摘录表','target_field':'原文摘录 + G-N八段式（existing row audit）',
          'normalized_conditions':normalize_conditions(r,'damage_mechanism'),'evidence_level':a.get('evidence_level','unreviewed'),
          'extraction_method':a.get('extraction_method'),'uncertainty':a.get('uncertainty'),
          'reject_reason':a.get('reject_reason'),'existing_record':True,'human_visual_review':False,
          'source_workbook':'02_损伤机理参考/条款摘录表_损伤机理.xlsx','source_row_candidate_id':a['candidate_id']})
    for a in rbi_audit:
        cid=a['clause_id']; r=rbi_rows.get(cid,{}); priority=rbi_priority.get(cid)
        quote=str(r.get('原文摘录') or a.get('原文摘录') or '')
        exact=a.get('exact_pdf_page_indexes')
        try: pidx=int(exact.split(';')[0]) if exact else None
        except (ValueError,AttributeError): pidx=None
        tableid=None
        m=re.search(r'(?:表|图)\s*([0-9A-Za-z.-]+)',quote)
        if m: tableid=m.group(1)
        out.append({'candidate_id':f"RBI-existing-{cid}",'standard_code':a['standard_code'],'clause_no':a['clause_no'],
          'printed_page':int(a['page']) if str(a.get('page','')).isdigit() else None,'pdf_page_index':pidx,'source_bbox':None,
          'raw_quote':quote,'raw_quote_sha256':digest(quote),'quote_kind':'text','table_id':tableid,'row_id':None,'col_id':None,
          'target_sheet':'摘录表','target_field':'原文摘录（existing row audit）',
          'normalized_conditions':normalize_conditions(r,'rbi'),'evidence_level':'machine_located' if pidx is not None else 'unreviewed',
          'extraction_method':'existing_RBI_workbook+PDF page text normalized match (machine locator only)',
          'uncertainty':('heuristic table explanation pattern; not adjudicated false; cell/header/unit/boundary require visual review' if priority else 'no human visual verification; machine page match is not evidence of quote exactness'),
          'reject_reason':None,'existing_record':True,'priority_table_quote_review':bool(priority),
          'source_workbook':'01_RBI标准/条款摘录表_RBI.xlsx','source_workbook_row':int(a['excel_row']) if str(a.get('excel_row','')).isdigit() else None,
          'extractor_signature_present':bool(a.get('抽取人')),'reviewer_signature_present':bool(a.get('复核人'))})
    for a in c_rows:
        out.append({**a,'candidate_id':a['candidate_id'],'existing_record':False,'human_visual_review':False})
    OUT.parent.mkdir(parents=True,exist_ok=True)
    with OUT.open('w',encoding='utf-8',newline='\n') as f:
        for row in out: f.write(json.dumps(row,ensure_ascii=False,separators=(',',':'))+'\n')
    ids=[x['candidate_id'] for x in out]
    print(f"records={len(out)} unique_ids={len(set(ids))} dm={len(dm_screen)} rbi={len(rbi_audit)} rbi_priority={len(rbi_priority)} supplementary={len(c_rows)}")
    print(f"registry={OUT} bytes={OUT.stat().st_size}")
    if len(ids)!=len(set(ids)): raise SystemExit('duplicate candidate_id')
if __name__=='__main__': main()
